"""30 сквозных сценариев по работающему docker-compose: от входа до выгрузки, через nginx интерфейса.

В отличие от `bench.e2e`, здесь работают настоящие API, воркеры, Redis и MinIO, запросы идут тем же путём, что из
браузера. Запускайте на заглушке LLM (`devtools/llm_stub`): проверяются связки и правила продукта, а не качество
находок, и часть сценариев рассчитана на её ответы. Созданные пакеты и правило удаляются в конце.

    python -m bench.e2e_api

Адрес и пароль demo-пользователей — `E2E_BASE_URL` и `DEMO_PASSWORD`. Код выхода 1, если сценарий не прошёл.
"""

from __future__ import annotations

import io
import os
import sys
import tempfile
import threading
import time
import traceback
import uuid
import zipfile
from collections.abc import Callable
from pathlib import Path

import httpx
import openpyxl
import pymupdf
from docx import Document as Docx

from bench import samples
from bench.scangen import make_scan

BASE = os.environ.get("E2E_BASE_URL", "http://localhost:8080").rstrip("/")
PW = os.environ.get("DEMO_PASSWORD", "demo")
TMP = Path(tempfile.mkdtemp(prefix="e2e30-"))
TAG = f"e2e-{uuid.uuid4().hex[:6]}"
created = []
ctx = {}


def client(login=None):
    c = httpx.Client(base_url=BASE, timeout=120)
    if login:
        r = c.post("/api/auth/login", json={"login": login, "password": PW})
        assert r.status_code == 200, r.text
    return c


M = client("methodologist")
S = client("specialist")
L = client("lawyer")


def ck(cond, msg=""):
    if not cond:
        raise AssertionError(msg)


def eq(r, code):
    ck(r.status_code == code, f"{r.request.method} {r.request.url.path}: {r.status_code}, ждали {code}: {r.text[:300]}")
    return r


def draft():
    b = eq(M.post("/api/batches"), 201).json()
    created.append(b["id"])
    return b["id"]


def upload(bid, *paths):
    files = [("files", (p.name, p.read_bytes())) for p in paths]
    return eq(M.post(f"/api/batches/{bid}/files", files=files), 201).json()


def start(bid, name):
    return eq(M.post(f"/api/batches/{bid}/start", json={"name": f"{TAG} {name}"}), 202)


def wait(bid, limit=300, done=("done", "partial_error", "error")):
    t = time.time()
    while time.time() - t < limit:
        v = eq(M.get(f"/api/batches/{bid}"), 200).json()
        if v["status"] in done:
            return v, time.time() - t
        time.sleep(1)
    raise AssertionError(f"пакет {bid} не закончился за {limit} с: {v['status']}, {[ (d['filename'], d['stage']) for d in v['documents']]}")


def run(name, *paths, limit=300):
    bid = draft()
    upload(bid, *paths)
    t = time.time()
    start(bid, name)
    v, _ = wait(bid, limit)
    return bid, v, time.time() - t


def doc(did, c=M):
    return eq(c.get(f"/api/documents/{did}"), 200).json()


def wait_export(eid):
    for _ in range(120):
        v = eq(M.get(f"/api/exports/{eid}"), 200).json()
        if v["status"] == "done":
            return eq(M.get(f"/api/exports/{eid}/download"), 200)
        ck(v["status"] != "error", f"выгрузка упала: {v}")
        time.sleep(1)
    raise AssertionError("выгрузка не готова за 120 с")


def fixtures():
    f = {}
    f["contract"] = samples.make_contract_docx(TMP / "Контракт.docx")
    f["notice"] = samples.make_notice_docx(TMP / "Извещение.docx")
    f["pdf"] = samples.make_contract_pdf(TMP / "Контракт_текст.pdf")
    f["pdf60"] = samples.make_contract_pdf(TMP / "Большой_60стр.pdf", items=samples.long_contract(60))
    src20 = samples.make_contract_pdf(TMP / "src20.pdf", items=samples.long_contract(20))
    f["scan"] = make_scan(src20, TMP / "Скан_20стр.pdf")
    (TMP / "обрезан.pdf").write_bytes(pymupdf.open(f["pdf60"]).tobytes()[:2000])
    f["cut"] = TMP / "обрезан.pdf"
    (TMP / "картинка.exe").write_bytes(b"MZ" + b"0" * 200)
    f["exe"] = TMP / "картинка.exe"
    d = pymupdf.open()
    d.new_page().insert_text((50, 50), "secret")
    d.save(TMP / "пароль.pdf", encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="1", owner_pw="2")
    f["enc"] = TMP / "пароль.pdf"
    with zipfile.ZipFile(TMP / "Пакет.zip", "w") as z:
        z.write(f["contract"], "Проект контракта/Контракт.docx")
        z.write(f["notice"], "Извещение/Извещение.docx")
    f["zip"] = TMP / "Пакет.zip"
    (TMP / "битый.zip").write_bytes(b"PK\x03\x04" + b"\x00" * 300)
    with zipfile.ZipFile(TMP / "Смешанный.zip", "w") as z:
        z.write(TMP / "обрезан.pdf", "Папка/сломан.pdf")
        z.writestr("Папка/мусор.docx", b"not a docx at all" * 20)
    f["mixzip"] = TMP / "Смешанный.zip"
    f["badzip"] = TMP / "битый.zip"
    w = Docx()
    w.add_paragraph("КОНТРАКТ на поставку товара")
    w.add_paragraph("1. Предмет контракта")
    w.add_paragraph("Цена Контракта составляет 100 000 (двести тысяч) рублей 00 копеек.")
    w.save(TMP / "Пропись.docx")
    f["words"] = TMP / "Пропись.docx"
    return f


F = fixtures()
cases: list[tuple[str, Callable]] = []


def case(title):
    def deco(fn):
        cases.append((title, fn))
        return fn
    return deco


# ---------- Доступ
@case("Служебные: healthz, config, SPA и глубокая ссылка через nginx")
def c01():
    ck(eq(httpx.get(BASE + "/api/healthz"), 200).json()["status"] == "ok")
    eq(M.get("/api/config"), 200)
    for p in ("/", "/batches/" + str(uuid.uuid4())):
        r = eq(httpx.get(BASE + p), 200)
        ck('id="root"' in r.text, "не SPA: " + p)


@case("Без входа: 401 в едином формате ошибок на русском")
def c02():
    for p in ("/api/me", "/api/batches", "/api/rules", f"/api/findings/{uuid.uuid4()}/source"):
        r = eq(httpx.get(BASE + p), 401)
        e = r.json()["error"]
        ck(e["code"] and any("а" <= ch <= "я" for ch in e["message"].lower()), r.text)


@case("Вход: неверный пароль 401; верный — httpOnly cookie и /me")
def c03():
    eq(httpx.post(BASE + "/api/auth/login", json={"login": "lawyer", "password": "wrong"}), 401)
    r = eq(httpx.post(BASE + "/api/auth/login", json={"login": "lawyer", "password": PW}), 200)
    ck("httponly" in r.headers["set-cookie"].lower(), r.headers["set-cookie"])
    ck(eq(L.get("/api/me"), 200).json()["role"] == "lawyer")


@case("Выход: cookie снята, сессия больше не работает")
def c04():
    c = client("lawyer")
    eq(c.post("/api/auth/logout"), 204)
    eq(c.get("/api/me"), 401)


@case("Роли: специалисту запрещены правила и черновик правила (403)")
def c05():
    eq(S.get("/api/rules"), 403)
    eq(S.post("/api/rules/draft", json={"description": "Заказчик требует страховать товар"}), 403)
    eq(M.get("/api/rules"), 200)


# ---------- Загрузка
@case("Загрузка DOCX в черновик и удаление файла")
def c06():
    bid = draft()
    v = upload(bid, F["contract"], F["notice"])
    ck(len(v["files"]) == 2 and all(x["check"] == "ok" for x in v["files"]), v)
    ck(v["files"][0]["filename"].endswith(".docx") and "Контракт" in {x["filename"].split(".")[0] for x in v["files"]}, "имя файла потеряно")
    v2 = eq(M.delete(f"/api/batches/{bid}/files/{v['files'][0]['id']}"), 200).json()
    ck(len(v2["files"]) == 1, v2)


@case("Обрезанный PDF: статус «повреждён», не 500")
def c07():
    bid = draft()
    v = upload(bid, F["cut"])
    ck(v["files"][0]["check"] == "corrupted", v)


@case("Неподдерживаемый и запароленный файлы получают свою причину")
def c08():
    bid = draft()
    v = upload(bid, F["exe"], F["enc"])
    by = {x["filename"]: x["check"] for x in v["files"]}
    ck(by == {"картинка.exe": "unsupported", "пароль.pdf": "encrypted"}, by)


@case("Запуск без пригодных файлов отклонён, пакет остаётся черновиком")
def c09():
    bid = draft()
    r = M.post(f"/api/batches/{bid}/start", json={"name": "пусто"})
    ck(r.status_code in (409, 422), f"{r.status_code} {r.text}")
    upload(bid, F["cut"])
    r = M.post(f"/api/batches/{bid}/start", json={"name": "только битый"})
    ck(r.status_code in (409, 422), f"битый файл: {r.status_code} {r.text}")
    ck(eq(M.get(f"/api/batches/{bid}"), 200).json()["status"] == "draft")


@case("Запуск без названия — ошибка поля; файлы в запущенный пакет не добавить (409)")
def c10():
    bid = draft()
    upload(bid, F["words"])
    r = M.post(f"/api/batches/{bid}/start", json={"name": "  "})
    ck(r.status_code == 422, f"{r.status_code} {r.text}")
    start(bid, "для 409")
    r = M.post(f"/api/batches/{bid}/files", files=[("files", ("x.docx", F["words"].read_bytes()))])
    eq(r, 409)
    wait(bid)


# ---------- Конвейер
@case("DOCX-пакет доходит до конца: находки, цвет, счётчики, SSE-события")
def c11():
    bid = draft()
    upload(bid, F["contract"], F["notice"])
    events = []
    def listen():
        with M.stream("GET", f"/api/batches/{bid}/events", timeout=120) as r:
            ck(r.headers["content-type"].startswith("text/event-stream"))
            for line in r.iter_lines():
                if line.startswith("event:"):
                    events.append(line.split(":", 1)[1].strip())
                    if events[-1] == "batch.done":
                        return
    th = threading.Thread(target=listen, daemon=True)
    th.start()
    time.sleep(0.5)
    start(bid, "docx")
    v, _ = wait(bid)
    th.join(30)
    ck(v["status"] == "done", v["status"])
    ck(v["counters"]["total"] == 2 and v["counters"]["errors"] == 0, v["counters"])
    ck(v["worst_level"] in ("critical", "warning"), v["worst_level"])
    ck(sum(v["findings"].values()) > 0, v["findings"])
    for need in ("batch.snapshot", "document.done", "batch.done"):
        ck(need in events, f"нет {need}: {sorted(set(events))}")
    ctx["docx"] = (bid, v)


@case("Каждая находка: цитата по смещениям существует в блоке, есть раздел или страница")
def c12():
    bid, v = ctx["docx"]
    n = 0
    for d in v["documents"]:
        view = doc(d["id"])
        blocks = {b["id"]: b for b in view["blocks"]}
        for f in view["findings"]:
            n += 1
            t = blocks[f["block_id"]]["text"]
            ck(0 <= f["start"] < f["end"] <= len(t), f)
            ck(t[f["start"]:f["end"]] == f["quote"], (t[f["start"]:f["end"]], f["quote"]))
            ck(f["section_number"] or f["section_title"] or f["page"], f"нет раздела: {f['id']}")
            ck(f["explanation"], "нет пояснения")
    ck(n > 0, "нет находок")


@case("PDF с текстовым слоем: страницы, находки, доступна сверка с оригиналом")
def c13():
    bid, v, _ = run("pdf", F["pdf"])
    ck(v["status"] == "done", v)
    d = v["documents"][0]
    ck(d["kind"] == "pdf" and d["pages"] >= 1, d)
    view = doc(d["id"])
    ck(view["findings"] and all(f["source_available"] for f in view["findings"]), "source_available")
    ctx["pdf"] = (bid, view)


@case("Скан на 20 страниц: быстрее 2 минут, уверенность OCR ≥ 0,95, пометка «со скана»")
def c14():
    bid, v, sec = run("скан", F["scan"], limit=240)
    ck(v["status"] == "done", [(d["filename"], d["error"]) for d in v["documents"]])
    d = v["documents"][0]
    ck(d["kind"] == "scan", d["kind"])
    ck(sec < 120, f"{sec:.0f} с")
    ck(d["ocr_confidence"] and d["ocr_confidence"] >= 0.95, d["ocr_confidence"])
    view = doc(d["id"])
    ck(view["findings"] and all(f["from_scan"] for f in view["findings"]), "from_scan")
    ctx["scan"] = (bid, view, sec)
    return f"{sec:.0f} с, уверенность {d['ocr_confidence']:.3f}"


@case("PDF на 60 страниц: загрузка и старт отвечают сразу, без 504")
def c15():
    bid = draft()
    t = time.time()
    upload(bid, F["pdf60"])
    up = time.time() - t
    t = time.time()
    start(bid, "60 страниц")
    st = time.time() - t
    v, sec = wait(bid, 400)
    ck(st < 3, f"старт {st:.1f} с")
    ck(v["status"] == "done", v["documents"][0]["error"])
    ck(v["documents"][0]["pages"] >= 50, v["documents"][0]["pages"])
    return f"{v['documents'][0]['pages']} стр., загрузка {up:.1f} с, старт {st:.2f} с, обработка {sec:.0f} с"


@case("Zip с папками и русскими именами: распакован, путь сохранён")
def c16():
    bid, v, _ = run("zip", F["zip"])
    ck(v["status"] == "done", [(d["filename"], d["status"], d["error"]) for d in v["documents"]])
    paths = {d.get("original_path") for d in v["documents"] if d["kind"] != "zip"}
    ck("Проект контракта/Контракт.docx" in paths and "Извещение/Извещение.docx" in paths, paths)
    ctx["zip"] = (bid, v)


@case("Битые файлы внутри архива рядом с хорошим документом: пакет доходит до конца, ошибка на русском, повтор упавшего")
def c17():
    bid = draft()
    v = upload(bid, F["badzip"])
    ck(v["files"][0]["check"] == "corrupted", v)
    eq(M.delete(f"/api/batches/{bid}/files/{v['files'][0]['id']}"), 200)
    upload(bid, F["mixzip"], F["words"])
    start(bid, "смешанный архив")
    v, _ = wait(bid)
    states = {d["filename"]: (d["status"], (d["error"] or {}).get("message")) for d in v["documents"]}
    good = next(d for d in v["documents"] if d["filename"] == "Пропись.docx")
    ck(good["status"] == "done", states)
    bad = [d for d in v["documents"] if d["status"] == "error"]
    ck(bad and v["status"] == "partial_error", states)
    ck(all(any("а" <= ch <= "я" for ch in d["error"]["message"].lower()) for d in bad), states)
    ck(v["counters"]["errors"] == len(bad), v["counters"])
    eq(M.post(f"/api/documents/{bad[0]['id']}/retry"), 202)
    v2, _ = wait(bid)
    ck(v2["status"] == "partial_error", v2["status"])
    ck(next(d for d in v2["documents"] if d["id"] == bad[0]["id"])["status"] == "error", "повтор битого файла")
    return str(states)


@case("Паспорт и финансовая нагрузка в meta пакета; повтор здорового документа — 409")
def c18():
    bid, v = ctx["docx"]
    meta = v["meta"]
    keys = [f["key"] for f in meta["passport"]["fields"]]
    ck(keys, "паспорт пуст")
    fin = meta.get("finance")
    ck(fin, f"нет finance при полях паспорта {keys}")
    ck(fin["securities"] or fin["cash_cycle"] or fin["deadline_at"], fin)
    eq(M.post(f"/api/documents/{v['documents'][0]['id']}/retry"), 409)
    sec = {s["key"]: (s["value"], s["amount"], s["basis"]) for s in fin["securities"]}
    ck(fin["total_security"] == round(sum(s["amount"] or 0 for s in fin["securities"]), 2), fin)
    return f"НМЦК {fin['nmck']}; {sec}; всего {fin['total_security']}; цикл {fin['cash_cycle']}; неполная {fin['total_incomplete']}"


@case("Проверка без LLM: цифры и пропись A-01; rule.id = null, фильтр check:A-01")
def c19():
    bid, v, _ = run("пропись", F["words"])
    ck(v["status"] == "done", v)
    view = doc(v["documents"][0]["id"])
    a = [f for f in view["findings"] if f.get("check_code") == "A-01"]
    ck(a, [f.get("check_code") or f["rule"]["code"] for f in view["findings"]])
    ck(a[0]["rule"]["id"] is None and a[0]["rule"]["code"] == "A-01", a[0]["rule"])
    r = eq(M.get(f"/api/batches/{bid}/findings", params={"rule_id": "check:A-01"}), 200).json()
    ck(r["items"] and all(f["check_code"] == "A-01" for f in r["items"]), r)
    ctx["words"] = (bid, v, view)


# ---------- Ревью
@case("Ревью: отклонение всех находок пересчитывает цвет до «Без рисков», возврат — обратно")
def c20():
    bid, v, view = ctx["words"]
    last = None
    start_level = view["worst_level"]
    for f in view["findings"]:
        last = eq(M.patch(f"/api/findings/{f['id']}", json={"version": f["version"], "status": "rejected"}), 200).json()
    ck(last["document"]["worst_level"] == "ok", last["document"])
    ck(last["batch"]["worst_level"] == "ok", last["batch"])
    ck(eq(M.get(f"/api/batches/{bid}"), 200).json()["worst_level"] == "ok")
    for f in doc(view["id"])["findings"]:
        last = eq(M.patch(f"/api/findings/{f['id']}", json={"version": f["version"], "status": "new"}), 200).json()
    ck(last["document"]["worst_level"] == start_level, (last["document"], start_level))


@case("Цвет — худший уровень: снижение критичной до «Внимание» меняет цвет документа")
def c21():
    bid, v = ctx["zip"]
    d = next(x for x in v["documents"] if x["worst_level"] == "critical")
    view = doc(d["id"])
    last = None
    for f in [f for f in view["findings"] if f["level"] == "critical"]:
        last = eq(M.patch(f"/api/findings/{f['id']}", json={"version": f["version"], "status": "edited", "level": "warning"}), 200).json()
    ck(last["document"]["worst_level"] == "warning", last["document"])
    f = doc(d["id"])["findings"][0]
    ck(f["original_level"] in ("critical", "warning"))


@case("Конфликт версий: устаревшая версия находки — 409 с актуальной версией")
def c22():
    _, view = ctx["pdf"]
    f = view["findings"][0]
    eq(M.patch(f"/api/findings/{f['id']}", json={"version": f["version"], "status": "accepted"}), 200)
    r = eq(L.patch(f"/api/findings/{f['id']}", json={"version": f["version"], "status": "rejected"}), 409)
    ck(r.json()["error"]["code"] == "version_conflict" and r.json()["error"]["details"]["current"], r.text)


@case("Смена уровня и возврат в «новая»: исходный уровень восстановлен, история записана")
def c23():
    _, view = ctx["pdf"]
    f = doc(view["id"])["findings"][-1]
    other = "warning" if f["level"] == "critical" else "critical"
    r = eq(M.patch(f"/api/findings/{f['id']}", json={"version": f["version"], "level": other, "status": "edited", "comment": "уточнил"}), 200).json()["finding"]
    ck(r["level"] == other and r["original_level"] == f["level"], r)
    r2 = eq(M.patch(f"/api/findings/{f['id']}", json={"version": r["version"], "status": "new"}), 200).json()["finding"]
    ck(r2["level"] == f["level"], r2)
    acts = [e["action"] for e in eq(M.get(f"/api/findings/{f['id']}/history"), 200).json()["events"]]
    ck("edited" in acts and "reopened" in acts, acts)


@case("Ручная пометка: выглядит как находка, автор — пользователь; удаление пересчитывает")
def c24():
    _, view = ctx["pdf"]
    b = next(x for x in view["blocks"] if len(x["text"]) > 30)
    r = eq(L.post(f"/api/documents/{view['id']}/findings", json={"block_id": b["id"], "start": 0, "end": 20, "level": "warning", "comment": "ручная"}), 201).json()
    f = r["finding"]
    ck(f["source"] == "user" and f["author"]["full_name"] and f["quote"] == b["text"][:20] and (f["section_number"] or f["section_title"] or f["page"]), f)
    eq(L.delete(f"/api/findings/{f['id']}"), 200)
    ck(all(x["id"] != f["id"] for x in doc(view["id"])["findings"]))
    eq(L.post(f"/api/documents/{view['id']}/findings", json={"block_id": b["id"], "start": 5, "end": 100000, "level": "warning"}), 422)


@case("Поиск по документу без учёта регистра и «ё»; блоки по диапазону страниц; % — обычный символ")
def c25():
    _, view = ctx["pdf"]
    r = eq(M.get(f"/api/documents/{view['id']}/search", params={"q": "ПОСТАВЩИК"}), 200).json()
    ck(r["matches"], r)
    r2 = eq(M.get(f"/api/documents/{view['id']}/search", params={"q": "%%%нет такого%%%"}), 200).json()
    ck(not r2["matches"], r2)
    eq(M.get(f"/api/documents/{view['id']}/blocks", params={"from_page": 1, "to_page": 1}), 200)
    lst = eq(M.get("/api/batches", params={"search": "%", "period": "all"}), 200).json()
    ck(all("%" in x["name"] for x in lst["items"]), "% сработал как шаблон LIKE")
    found = eq(M.get("/api/batches", params={"search": TAG}), 200).json()
    ck(found["items"], "поиск пакета по имени")


# ---------- Сверка с оригиналом
@case("Сверка с оригиналом, PDF: рамки в долях страницы, JPEG, нет страницы — 404")
def c26():
    _, view = ctx["pdf"]
    f = doc(view["id"])["findings"][0]
    r = eq(M.get(f"/api/findings/{f['id']}/source"), 200).json()
    ck(r["rects"], f"нет рамок: {r}")
    ck(all(0 <= x["x0"] < x["x1"] <= 1 and 0 <= x["y0"] < x["y1"] <= 1 for x in r["rects"]), r["rects"])
    img = eq(M.get(r["image_url"]), 200)
    ck(img.content[:2] == b"\xff\xd8" and img.headers["content-type"] == "image/jpeg")
    eq(M.get(f"/api/documents/{view['id']}/pages/9999/image"), 404)
    eq(httpx.get(BASE + r["image_url"]), 401)
    eq(M.get(f"/api/documents/{view['id']}/pages/0/image"), 422)


@case("Сверка с оригиналом: скан — рамки OCR; DOCX — кнопки нет и 409")
def c27():
    _, view, _ = ctx["scan"]
    f = view["findings"][0]
    r = eq(M.get(f"/api/findings/{f['id']}/source"), 200).json()
    ck(r["rects"], f"нет рамок на скане: {r}")
    eq(M.get(r["image_url"]), 200)
    _, v = ctx["docx"]
    dv = doc(v["documents"][0]["id"])
    ck(not dv["findings"][0]["source_available"], "у DOCX source_available")
    eq(M.get(f"/api/findings/{dv['findings'][0]['id']}/source"), 409)


# ---------- Правила
@case("Новое правило без деплоя: тест, сохранение, срабатывает только в новых пакетах; review в списке")
def c28():
    rule = {"code": "R-77", "title": "Страхование товара поставщиком", "description": "Поставщик страхует товар за свой счёт",
            "level": "critical", "kind": "semantic_llm",
            "prompt_template": "Фрагмент договора: {фрагмент}. Риск, если поставщик обязан застраховать товар за свой счёт."}
    text = "Страхование товара поставщиком: Поставщик обязан застраховать товар за свой счёт на всю стоимость."
    tested = eq(M.post("/api/rules/test", json={"rule": rule, "text": text}), 200).json()
    ck(tested["findings"], f"тест правила на своём тексте ничего не нашёл: {tested}")
    eq(M.post("/api/rules", json={**rule, "prompt_template": "без переменной"}), 422)
    saved = eq(M.post("/api/rules", json=rule), 201).json()
    try:
        w = Docx()
        w.add_paragraph("1. Страхование")
        w.add_paragraph("Поставщик обязан застраховать товар за свой счёт на всю стоимость страхование товара поставщиком.")
        w.save(TMP / "Страхование.docx")
        old_bid, old = ctx["docx"]
        ck(not any(t["code"] == "R-77" for t in old.get("rules_triggered") or []), "старый пакет")
        bid, v, _ = run("новое правило", TMP / "Страхование.docx")
        codes = {t["code"] for t in v["rules_triggered"]}
        ck("R-77" in codes, codes)
        ck(not any(t["code"] == "R-77" for t in eq(M.get(f"/api/batches/{old_bid}"), 200).json()["rules_triggered"]), "правило задело старый пакет")
        rows = eq(M.get("/api/rules", params={"search": "R-77"}), 200).json()
        ck(rows[0]["review"] == {"confirmed": 0, "rejected": 0, "precision": None, "noisy": False}, rows[0])
        d = doc(v["documents"][0]["id"])
        f = next(x for x in d["findings"] if x["rule"]["code"] == "R-77")
        eq(M.patch(f"/api/findings/{f['id']}", json={"version": f["version"], "status": "rejected"}), 200)
        row = eq(M.get("/api/rules", params={"search": "R-77"}), 200).json()[0]
        ck(row["review"]["rejected"] == 1 and row["review"]["precision"] == 0.0, row)
    finally:
        eq(M.delete(f"/api/rules/{saved['id']}"), 204)


# ---------- Выгрузка
@case("Выгрузка документа: DOCX с комментариями Word, PDF с выделениями")
def c29():
    _, v = ctx["docx"]
    d = next(x for x in v["documents"] if sum(x["findings"].values()) > 0)
    r = wait_export(eq(M.post(f"/api/documents/{d['id']}/exports", json={}), 202).json()["export_id"])
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        ck("word/comments.xml" in z.namelist(), z.namelist())
        ck(b"w:comment " in z.read("word/comments.xml"))
    ck("attachment" in r.headers["content-disposition"])
    _, view = ctx["pdf"]
    r = wait_export(eq(M.post(f"/api/documents/{view['id']}/exports", json={}), 202).json()["export_id"])
    with pymupdf.open(stream=r.content, filetype="pdf") as p:
        ck(sum(len(list(pg.annots() or [])) for pg in p) > 0, "нет аннотаций")


@case("Выгрузка пакета: структура архива, Сводка.xlsx с паспортом и суммой обеспечений; ask, вывод, запрос, черновик; удаление")
def c30():
    bid, v = ctx["zip"]
    r = wait_export(eq(M.post(f"/api/batches/{bid}/exports", json={}), 202).json()["export_id"])
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        names = z.namelist()
        ck("Сводка.xlsx" in names, names)
        ck(any(n.startswith("Проект контракта/") for n in names) and any(n.startswith("Извещение/") for n in names), names)
        wb = openpyxl.load_workbook(io.BytesIO(z.read("Сводка.xlsx")))
    rows = [str(row[0]) for ws in wb for row in ws.iter_rows(values_only=True) if row and row[0]]
    ck("Паспорт закупки" in rows, rows[:30])
    fin = eq(M.get(f"/api/batches/{bid}"), 200).json()["meta"].get("finance") or {}
    if fin.get("total_security") is not None:
        ck("Обеспечения всего, руб." in rows, rows[:40])
    a = eq(M.post(f"/api/batches/{bid}/ask", json={"question": "Какой срок поставки?"}), 200).json()
    ck("found" in a, a)
    eq(M.post(f"/api/batches/{bid}/ask", json={"question": "  "}), 422)
    b = eq(M.post(f"/api/batches/{bid}/brief"), 200).json()
    ck(b, b)
    ck(eq(M.get(f"/api/batches/{bid}"), 200).json()["meta"].get("brief"), "вывод не сохранён")
    c = wait_export(eq(M.post(f"/api/batches/{bid}/clarification"), 202).json()["export_id"])
    ck(any(p.text.strip() for p in Docx(io.BytesIO(c.content)).paragraphs), "пустой запрос")
    dr = eq(M.post("/api/rules/draft", json={"description": "Заказчик требует страховать товар за счёт поставщика"}), 200).json()
    ck(dr["code"].startswith("R-"), dr)
    x = draft()
    upload(x, F["words"])
    start(x, "удаление")
    wait(x)
    eq(M.delete(f"/api/batches/{x}"), 204)
    eq(M.get(f"/api/batches/{x}"), 404)


def main() -> int:
    failed = 0
    for i, (title, fn) in enumerate(cases, 1):
        began = time.time()
        try:
            note, status = fn() or "", "PASS"
        except Exception as exc:  # noqa: BLE001 — падение сценария не останавливает остальные
            line = traceback.extract_tb(exc.__traceback__)[-1].lineno
            note, status = f"{type(exc).__name__}: {str(exc)[:500]} (строка {line})", "FAIL"
            failed += 1
        print(f"{i:>2} {status} {time.time() - began:6.1f}s {title}" + (f"\n     {note}" if note else ""), flush=True)
    for bid in created:
        M.delete(f"/api/batches/{bid}")
    print(f"\nИТОГО: {len(cases) - failed} из {len(cases)}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
