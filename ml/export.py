"""Разметка исходного документа находками (FR-7.1, FR-7.2): подсветка и комментарии.

Задача экспорта (очередь `export`) берёт находки из `db.exports.export_payload` и исходный файл
из MinIO, а смещения находки переводит в исходник по карте позиций блока (`source_map`, контракт 2):

- DOCX (и DOC, RTF — через сконвертированный DOCX): runs, в которые попала цитата, режутся по её
  границам, кусок внутри цитаты получает заливку уровня, а вся цитата — комментарий Word с правилом,
  объяснением и рекомендацией;
- PDF и сканы: аннотация-подсветка по рамкам слов цитаты (у скана — рамки слов OCR) со всплывающей
  заметкой того же текста.

Номер абзаца и run считаются так же, как их считает парсер: `part.element.iter(w:p)` и
`paragraph.iter(w:r)`, смещение внутри run — по тем же дочерним элементам (`w:t`, `w:tab`, `w:br`,
…), поэтому подсветка ложится ровно на символы `block.text[start:end]`.
"""

from __future__ import annotations

import copy
import logging
from collections import defaultdict
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from contracts.extraction import BBox, SourceMap

log = logging.getLogger(__name__)

AUTHOR = "Светофор рисков"
INITIALS = "СР"
LEVEL_LABELS = {"critical": "Критично", "warning": "Внимание"}
STATUS_LABELS = {"new": "Новая", "accepted": "Принята", "rejected": "Отклонена", "edited": "Изменена"}
# Заливка DOCX и цвет подсветки PDF: светлые тона, текст под ними читается
DOCX_FILL = {"critical": "FFC7CE", "warning": "FFEB9C", "rejected": "D9D9D9"}
PDF_COLOR = {"critical": (1.0, 0.55, 0.55), "warning": (1.0, 0.85, 0.3), "rejected": (0.8, 0.8, 0.8)}
# Чем выше, тем позже кладётся заливка: на пересечении находок виден худший уровень
_PAINT_ORDER = {"rejected": 0, "warning": 1, "critical": 2}

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_FOOTNOTES_PART = "/word/footnotes.xml"


def _w(tag: str) -> str:
    return f"{{{_W}}}{tag}"


# Порядок дочерних элементов `w:rPr` по схеме: Word не открывает файл, где `w:shd` стоит не на месте
_RPR_ORDER = [_w(t) for t in (
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
    "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
    "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
    "em", "lang", "eastAsianLayout", "specVanish", "oMath",
)]


@dataclass
class ExportStats:
    """Итог разметки: сколько находок легло на исходник и какие не удалось показать."""

    findings: int = 0
    placed: int = 0
    comments: int = 0
    missed: list[str] = field(default_factory=list)


# --- Текст комментария -----------------------------------------------------------------------


def _get(f: Any, key: str, default: Any = None) -> Any:
    if isinstance(f, Mapping):
        return f.get(key, default)
    return getattr(f, key, default)


def _paint_key(f: Any) -> str:
    if _get(f, "status") == "rejected":
        return "rejected"
    return "critical" if _get(f, "level") == "critical" else "warning"


def finding_note(f: Any) -> list[str]:
    """Абзацы комментария к находке: заголовок, место, объяснение, рекомендация, решение юриста."""
    level = LEVEL_LABELS.get(_get(f, "level"), "Внимание")
    rule = _get(f, "rule") or {}
    code, title = _get(rule, "code"), _get(rule, "title")
    if code or title:
        head = f"{level} · {' '.join(x for x in (code, title) if x)}"
    else:
        author = _get(_get(f, "author") or {}, "full_name")
        head = f"{level} · ручная пометка" + (f" ({author})" if author else "")
    lines = [head]
    place = []
    if _get(f, "section_number"):
        place.append(f"п. {_get(f, 'section_number')}")
    if _get(f, "section_title"):
        place.append(str(_get(f, "section_title")))
    if _get(f, "page"):
        place.append(f"стр. {_get(f, 'page')}")
    if place:
        lines.append("Раздел: " + ", ".join(place))
    explanation = (_get(f, "explanation") or "").strip()
    if explanation:
        lines.append(explanation)
    recommendation = (_get(f, "recommendation") or "").strip()
    if recommendation:
        lines.append(f"Рекомендация: {recommendation}")
    status = _get(f, "status")
    if status and status != "new":
        decided = _get(_get(f, "decided_by") or {}, "full_name")
        lines.append(f"Статус: {STATUS_LABELS.get(status, status)}" + (f" ({decided})" if decided else ""))
    comment = (_get(f, "comment") or "").strip()
    if comment and comment != explanation:
        lines.append(f"Комментарий: {comment}")
    return lines


def _source_map(f: Any) -> SourceMap | None:
    raw = _get(f, "source_map")
    if raw is None or isinstance(raw, SourceMap):
        return raw
    return SourceMap.model_validate(raw)


# --- DOCX ------------------------------------------------------------------------------------


def _run_items(r) -> list[tuple[Any, int]]:
    """Дочерние элементы run и их длина в `Run.text` — как их считает парсер."""
    out = []
    for child in r:
        tag = child.tag
        if tag == _w("t"):
            n = len(child.text or "")
        elif tag in (_w("tab"), _w("ptab"), _w("cr"), _w("noBreakHyphen")):
            n = 1
        elif tag == _w("br"):
            n = 1 if child.get(_w("type")) in (None, "textWrapping") else 0
        else:
            n = 0
        out.append((child, n))
    return out


def _split_run(r, k: int):
    """Делит run на `[0, k)` и `[k, …)`: правая часть — новый run с теми же свойствами сразу после."""
    pos = 0
    items = _run_items(r)
    cut = None  # индекс первого элемента правой части
    for i, (child, n) in enumerate(items):
        if child.tag == _w("rPr"):
            continue
        if pos >= k:
            cut = i
            break
        if pos + n > k:
            # Разрез внутри w:t: текст делится на два w:t
            text = child.text or ""
            left, right = text[: k - pos], text[k - pos:]
            child.text = left
            child.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            twin = copy.deepcopy(child)
            twin.text = right
            child.addnext(twin)
            cut = i + 1
            break
        pos += n
    right_run = copy.deepcopy(r)
    for child in list(right_run):
        if child.tag != _w("rPr"):
            right_run.remove(child)
    if cut is not None:
        for child in list(r)[cut:]:
            right_run.append(child)  # append переносит элемент из исходного run
    r.addnext(right_run)
    return right_run


def _shade(r, fill: str) -> None:
    rpr = r.find(_w("rPr"))
    if rpr is None:
        rpr = r.makeelement(_w("rPr"), {})
        r.insert(0, rpr)
    old = rpr.find(_w("shd"))
    if old is not None:
        rpr.remove(old)
    shd = rpr.makeelement(_w("shd"), {_w("val"): "clear", _w("color"): "auto", _w("fill"): fill})
    rank = _RPR_ORDER.index(_w("shd"))
    after = next((c for c in rpr if c.tag in _RPR_ORDER and _RPR_ORDER.index(c.tag) > rank), None)
    if after is not None:
        after.addprevious(shd)
    else:
        rpr.append(shd)


@dataclass
class _Piece:
    run: Any
    start: int  # смещение в тексте run
    end: int


def _resolve(sm: SourceMap, start: int, end: int, paragraphs: Sequence, run_cache: dict) -> list[_Piece]:
    pieces = []
    for span in sm.runs:
        s, e = max(start, span.start), min(end, span.end)
        if s >= e:
            continue
        if not 0 <= span.paragraph_index < len(paragraphs):
            continue
        p = paragraphs[span.paragraph_index]
        runs = run_cache.get(p)
        if runs is None:
            runs = run_cache[p] = list(p.iter(_w("r")))
        if not 0 <= span.run_index < len(runs):
            continue
        offset = span.run_offset - span.start
        pieces.append(_Piece(runs[span.run_index], s + offset, e + offset))
    return pieces


def _cut_runs(pieces: Iterable[_Piece]) -> dict[int, list[tuple[int, int, Any]]]:
    """Режет каждый затронутый run по границам всех цитат сразу; возвращает отрезки run → элементы."""
    cuts: dict[int, set[int]] = defaultdict(set)
    runs: dict[int, Any] = {}
    for piece in pieces:
        runs[id(piece.run)] = piece.run
        cuts[id(piece.run)].update((piece.start, piece.end))
    segments: dict[int, list[tuple[int, int, Any]]] = {}
    for key, r in runs.items():
        total = sum(n for _, n in _run_items(r))
        points = sorted(c for c in cuts[key] if 0 < c < total)
        bounds = [0, *points, total]
        parts: list[Any] = [r]
        for k in reversed(points):
            parts.insert(1, _split_run(r, k))
        segments[key] = [(bounds[i], bounds[i + 1], parts[i]) for i in range(len(parts))]
    return segments


def _segment_runs(piece: _Piece, segments: dict[int, list[tuple[int, int, Any]]]) -> list[Any]:
    return [el for s, e, el in segments.get(id(piece.run), []) if s < piece.end and e > piece.start]


def annotate_docx(src: str | Path, findings: Sequence[Any], out: str | Path, *, author: str = AUTHOR) -> ExportStats:
    """DOCX с заливкой цитат и комментариями Word. `findings` — элементы `export_payload()["findings"]`.

    Для DOC и RTF передаётся сконвертированный DOCX (`document.normalized_key`): по нему считалась карта
    runs. Находки сносок получают заливку без комментария — Word не хранит комментарии в сносках.
    """
    from docx import Document
    from docx.text.run import Run
    from lxml import etree

    doc = Document(str(src))
    body_paragraphs = list(doc.element.iter(_w("p")))
    foot_part = next((p for p in doc.part.package.iter_parts() if str(p.partname) == _FOOTNOTES_PART), None)
    foot_root = etree.fromstring(foot_part.blob) if foot_part is not None else None
    foot_paragraphs = list(foot_root.iter(_w("p"))) if foot_root is not None else []

    stats = ExportStats(findings=len(findings))
    run_cache: dict = {}
    resolved: list[tuple[Any, list[_Piece], str]] = []
    for f in findings:
        sm = _source_map(f)
        start, end = int(_get(f, "start", 0)), int(_get(f, "end", 0))
        pieces: list[_Piece] = []
        part = "document"
        if sm is not None and sm.format == "docx" and end > start:
            part = sm.part
            paragraphs = foot_paragraphs if part == "footnotes" else body_paragraphs
            pieces = _resolve(sm, start, end, paragraphs, run_cache)
        if not pieces:
            stats.missed.append(str(_get(f, "id", "")))
            log.info("annotate_docx: находка %s не легла на runs исходника", _get(f, "id"))
            continue
        resolved.append((f, pieces, part))

    segments = _cut_runs(p for _, pieces, _ in resolved for p in pieces)
    for f, pieces, _part in sorted(resolved, key=lambda x: _PAINT_ORDER[_paint_key(x[0])]):
        fill = DOCX_FILL[_paint_key(f)]
        for piece in pieces:
            for r in _segment_runs(piece, segments):
                _shade(r, fill)
        stats.placed += 1

    # Комментарии — после всех разрезов: иначе метки диапазона сдвинули бы номера runs
    for f, pieces, part in resolved:
        if part != "document":
            continue
        runs = [r for piece in pieces for r in _segment_runs(piece, segments)]
        if not runs:
            continue
        lines = finding_note(f)
        parent = doc.paragraphs[0] if doc.paragraphs else None
        comment = doc.add_comment([Run(runs[0], parent), Run(runs[-1], parent)], text=lines[0], author=author,
                                  initials=INITIALS)
        comment.paragraphs[0].runs[0].bold = True
        for line in lines[1:]:
            comment.add_paragraph(line)
        stats.comments += 1

    if foot_part is not None and any(part == "footnotes" for _, _, part in resolved):
        foot_part._blob = etree.tostring(foot_root, xml_declaration=True, encoding="UTF-8", standalone=True)
    doc.save(str(out))
    return stats


# --- PDF -------------------------------------------------------------------------------------


def _line_rects(boxes: list[BBox]) -> list[tuple[float, float, float, float]]:
    """Рамки слов одной строки сливаются в одну: подсветка без щелей между словами."""
    lines: list[list[BBox]] = []
    for b in sorted(boxes, key=lambda b: (b.y0 + b.y1) / 2):
        mid = (b.y0 + b.y1) / 2
        if lines:
            top, bottom = min(x.y0 for x in lines[-1]), max(x.y1 for x in lines[-1])
            if top <= mid <= bottom:
                lines[-1].append(b)
                continue
        lines.append([b])
    rects: list[tuple[float, float, float, float]] = []
    for line in lines:
        height = max(max(b.y1 - b.y0 for b in line), 1.0)
        current: list[float] | None = None
        for b in sorted(line, key=lambda b: b.x0):
            if current is not None and b.x0 - current[2] < height * 1.5:
                current = [min(current[0], b.x0), min(current[1], b.y0), max(current[2], b.x1), max(current[3], b.y1)]
                rects[-1] = tuple(current)
                continue
            current = [b.x0, b.y0, b.x1, b.y1]
            rects.append(tuple(current))
    return rects


def annotate_pdf(src: str | Path, findings: Sequence[Any], out: str | Path, *, author: str = AUTHOR) -> ExportStats:
    """PDF с аннотациями-подсветками и заметками, видимыми в Acrobat и браузере (FR-7.2).

    Рамки берутся из слов цитаты (`source_map.words`, у скана — слова OCR), без слов — из рамки блока.
    PyMuPDF не работает в нескольких потоках: вызывать в воркере с пулом процессов.
    """
    import pymupdf

    stats = ExportStats(findings=len(findings))
    with pymupdf.open(str(src)) as doc:
        for f in sorted(findings, key=lambda x: _PAINT_ORDER[_paint_key(x)]):
            sm = _source_map(f)
            start, end = int(_get(f, "start", 0)), int(_get(f, "end", 0))
            by_page: dict[int, list[BBox]] = defaultdict(list)
            if sm is not None:
                for w in sm.words:
                    if w.start < end and w.end > start:
                        by_page[w.page].append(w.bbox)
            if not by_page and _get(f, "bbox") and _get(f, "page"):
                bbox = _get(f, "bbox")
                by_page[int(_get(f, "page"))].append(bbox if isinstance(bbox, BBox) else BBox.model_validate(bbox))
            pages = {p: boxes for p, boxes in by_page.items() if 1 <= p <= doc.page_count}
            if not pages:
                stats.missed.append(str(_get(f, "id", "")))
                log.info("annotate_pdf: находка %s без рамок на страницах документа", _get(f, "id"))
                continue
            lines = finding_note(f)
            color = PDF_COLOR[_paint_key(f)]
            for no, boxes in sorted(pages.items()):
                page = doc[no - 1]
                quads = [pymupdf.Rect(*r).quad for r in _line_rects(boxes)]
                annot = page.add_highlight_annot(quads=quads)
                if annot is None:
                    continue
                annot.set_colors(stroke=color)
                annot.set_info(title=author, subject=lines[0], content="\n".join(lines))
                annot.update()
            stats.placed += 1
            stats.comments += 1
        doc.save(str(out), garbage=3, deflate=True)
    return stats


def annotate(src: str | Path, findings: Sequence[Any], out: str | Path, *, author: str = AUTHOR) -> ExportStats:
    """Формат по расширению исходника: `.docx` — комментарии Word, `.pdf` — аннотации PDF."""
    suffix = Path(src).suffix.lower()
    if suffix == ".docx":
        return annotate_docx(src, findings, out, author=author)
    if suffix == ".pdf":
        return annotate_pdf(src, findings, out, author=author)
    raise ValueError(f"Разметка не поддерживает {suffix}: нужен DOCX или PDF")
