"""Черновик запроса о даче разъяснений положений извещения по находкам пакета.

Найденный риск поставщик чаще всего снимает запросом к заказчику до окончания подачи заявок.
Число запросов по одной закупке ограничено (по 44-ФЗ — не больше трёх от участника), поэтому
все вопросы собираются в один документ:

- цитата, документ, пункт и страница каждого вопроса берутся из находки, а не из ответа
  модели, — в запрос не попадает условие, которого нет в документации;
- LLM только формулирует вопрос деловым языком: просьба разъяснить условие и предложение,
  как его изменить, по пояснению и рекомендации находки;
- без LLM (или при её сбое) вопрос собирается из пояснения и рекомендации шаблоном.

Нормы в `LEGAL_BASIS` сверяются с действующей редакцией закона перед использованием.
"""

from __future__ import annotations

import logging
import re
from collections.abc import Sequence
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

from pydantic import BaseModel, Field

from contracts.analysis import LLMUsage
from ml.analysis.llm import LLMClient, default_client
from ml.config import get_settings
from ml.errors import LLMError

log = logging.getLogger(__name__)

LEGAL_BASIS = {
    "44": "частью 5 статьи 42 Федерального закона от 05.04.2013 № 44-ФЗ «О контрактной системе в сфере закупок "
          "товаров, работ, услуг для обеспечения государственных и муниципальных нужд»",
    "223": "частью 2 статьи 3.2 Федерального закона от 18.07.2011 № 223-ФЗ «О закупках товаров, работ, услуг "
           "отдельными видами юридических лиц»",
}
MAX_ITEMS = 30
GROUP = 6


class ClarificationItem(BaseModel):
    finding_id: str | None = None
    document: str | None = Field(None, description="Имя файла документа")
    section_number: str | None = None
    page: int | None = None
    quote: str
    request: str = Field(description="Вопрос и предложение заказчику деловым языком")


class ClarificationDraft(BaseModel):
    title: str
    addressee: str | None = None
    procurement: str | None = None
    intro: str
    items: list[ClarificationItem]
    closing: str
    usage: LLMUsage = Field(default_factory=LLMUsage)

    def as_text(self) -> str:
        """Текст для поля запроса на электронной площадке."""
        lines = [self.intro, ""]
        for i, item in enumerate(self.items, 1):
            lines += [f"{i}. {_where(item)} установлено: «{_quote(item)}».", item.request, ""]
        lines.append(self.closing)
        return "\n".join(lines).strip()


class _Request(BaseModel):
    index: int
    request: str


class _Answer(BaseModel):
    items: list[_Request]


SYSTEM = """Ты — юрист поставщика. Готовишь запрос о даче разъяснений положений извещения об осуществлении закупки: письмо участника закупки заказчику.

Для каждого пронумерованного условия напиши один пункт запроса (request): 2–3 предложения делового стиля. Сначала вежливо попроси разъяснить условие, затем предложи, как его изменить, опираясь на пояснение и рекомендацию. Пиши от лица участника закупки («Просим разъяснить…», «Предлагаем…», «Просим рассмотреть возможность…»). Не цитируй условие — цитата уже есть в письме. Не оценивай действия заказчика словами «незаконно», «нарушает», не угрожай жалобой; ссылку на норму закона давай, только если она есть в пояснении или рекомендации.

Ответ — только JSON по схеме: {"items": [{"index": 1, "request": "..."}]} — по одному элементу на каждое условие."""


def _quote(item: ClarificationItem) -> str:
    """Цитата без точки в конце: точка ставится после кавычки."""
    return item.quote.rstrip(" .")


def _where(item: ClarificationItem) -> str:
    parts = []
    if item.section_number:
        parts.append(f"В пункте {item.section_number}")
    elif item.page:
        parts.append(f"На странице {item.page}")
    else:
        parts.append("В документации")
    if item.document:
        parts.append(f"документа «{item.document}»")
    return " ".join(parts)


_REQUEST = re.compile(
    r"^(?:направить\s+(?:заказчику\s+)?запрос\s+(?:на|о)\s+(?:даче\s+)?разъяснени[еяй]\s*(?:положений\s+извещения)?"
    r"\s*(?:и\s+)?)?(?P<verb>предложить|попросить|просить|потребовать)\s+(?:заказчику\s+|у\s+заказчика\s+)?(?P<rest>.+)",
    re.IGNORECASE,
)


def _group_index(index: int, group: list[int]) -> int | None:
    """Номер условия из ответа модели → сквозной номер в пределах своей группы.

    Модель может вернуть сквозной номер («7» для первого условия второй группы) или номер внутри
    группы («1»). Номер чужой группы не принимается: иначе ответ затёр бы вопрос другого условия.
    """
    if index in group:
        return index
    if 1 <= index <= len(group):
        return group[index - 1]
    return None


def _fallback(finding: dict) -> str:
    """Вопрос без LLM: просьба разъяснить, пояснение риска и предложение из рекомендации.

    «…и предложить срок не менее 30 дней» → «Предлагаем срок не менее 30 дней»; рекомендации,
    которые поставщик выполняет сам («проверить лицензию»), в письмо не попадают.
    """
    out = "Просим разъяснить данное условие."
    text = (finding.get("explanation") or "").strip()
    if text:
        out += f" {text}"
    m = _REQUEST.match((finding.get("recommendation") or "").strip())
    if m:
        verb = "Предлагаем" if m.group("verb").lower() == "предложить" else "Просим"
        rest = m.group("rest").strip()
        out += f" {verb} {rest}" + ("" if rest.endswith(".") else ".")
    return out.strip()


def draft_clarification(
    findings: Sequence[dict],
    *,
    subject: str | None = None,
    customer: str | None = None,
    procurement_number: str | None = None,
    law: str = "44",
    llm: LLMClient | None = None,
    use_llm: bool = True,
) -> ClarificationDraft:
    """Черновик запроса по находкам: словари из `db.findings.findings_payload` (нужны `quote`,
    `explanation`, `recommendation`, `section_number`, `page`, `document` или `filename`).

    Отклонённые находки и ручные пометки без пояснения тоже можно передать — отбор делает
    вызывающий (по умолчанию — активные находки пакета).
    """
    chosen = [f for f in findings if (f.get("quote") or "").strip()][:MAX_ITEMS]
    items = [
        ClarificationItem(
            finding_id=f.get("id"),
            document=(f.get("document") or {}).get("filename") if isinstance(f.get("document"), dict)
            else f.get("filename"),
            section_number=f.get("section_number"), page=f.get("page"), quote=f["quote"].strip(),
            request=_fallback(f),
        )
        for f in chosen
    ]
    usage = LLMUsage()
    if items and use_llm and (llm is not None or get_settings().llm_model):
        client = llm or default_client()
        conditions = []
        for i, (item, f) in enumerate(zip(items, chosen, strict=True), 1):
            conditions.append((i, (
                f"[{i}] {_where(item)}: «{item.quote}»\n"
                f"Риск: {(f.get('rule') or {}).get('title') or 'ручная пометка'}. {f.get('explanation') or ''}\n"
                f"Рекомендация: {f.get('recommendation') or '—'}"
            )))
        # Группы по GROUP условий параллельно: на длинном списке модель отвечала только на первое
        groups = [conditions[i:i + GROUP] for i in range(0, len(conditions), GROUP)]

        def ask(group):
            indexes = ", ".join(str(i) for i, _ in group)
            messages = [{"role": "system", "content": SYSTEM}, {"role": "user", "content": (
                f"Предмет закупки: {subject or 'не указан'}\n\nУсловия:\n\n" + "\n\n".join(t for _, t in group)
                + f"\n\nВерни ровно {len(group)} элементов items — по одному на каждое условие, index: {indexes}.")}]
            return client.chat(messages, _Answer, schema_name="clarification_request",
                               max_tokens=200 + 220 * len(group), reasoning="none")

        with ThreadPoolExecutor(max_workers=min(len(groups), get_settings().llm_max_concurrency)) as pool:
            futures = [pool.submit(ask, g) for g in groups]
            for group, future in zip(groups, futures, strict=True):
                try:
                    answer = future.result()
                except LLMError as exc:
                    log.warning("clarification: LLM недоступна, вопросы группы собраны шаблоном: %s", exc)
                    continue
                usage.add(answer.usage)
                for req in answer.data.items:
                    index = _group_index(req.index, [i for i, _ in group])
                    if index is not None and req.request.strip():
                        items[index - 1].request = req.request.strip()

    basis = LEGAL_BASIS.get(law, LEGAL_BASIS["44"])
    about = f" № {procurement_number}" if procurement_number else ""
    subject_part = f" (объект закупки: {subject})" if subject else ""
    intro = (f"В соответствии с {basis} просим дать разъяснения положений извещения об осуществлении "
             f"закупки{about}{subject_part} и прилагаемых к нему документов по следующим вопросам.")
    closing = ("Просим разместить разъяснения в единой информационной системе в сфере закупок в установленный "
               "срок. Ответы позволят участникам закупки корректно подготовить заявки и исполнить контракт.")
    return ClarificationDraft(
        title="Запрос о даче разъяснений положений извещения об осуществлении закупки",
        addressee=customer,
        procurement=" — ".join(x for x in (f"Закупка № {procurement_number}" if procurement_number else None,
                                           subject) if x) or None,
        intro=intro, items=items, closing=closing, usage=usage,
    )


def render_clarification_docx(draft: ClarificationDraft, path: str | Path, *, sender: str | None = None,
                              on_date: date | None = None) -> Path:
    """Черновик запроса в DOCX: шапка, вступление, нумерованные вопросы с цитатами, подпись."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = Document()
    for section in doc.sections:
        section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    if draft.addressee:
        p = doc.add_paragraph(f"Заказчику: {draft.addressee}")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph(f"Дата: {(on_date or date.today()).strftime('%d.%m.%Y')}")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(draft.title.upper())
    run.bold = True
    if draft.procurement:
        sub = doc.add_paragraph(draft.procurement)
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph(draft.intro)
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for i, item in enumerate(draft.items, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(f"{i}. {_where(item)}").bold = True
        p.add_run(" установлено: ")
        p.add_run(f"«{_quote(item)}»").italic = True
        p.add_run(".")
        q = doc.add_paragraph(item.request)
        q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        q.paragraph_format.first_line_indent = Cm(1.0)

    closing = doc.add_paragraph(draft.closing)
    closing.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()
    doc.add_paragraph(f"Участник закупки: {sender or '________________________'}")
    doc.add_paragraph("Подпись ____________________    М. П.")
    doc.core_properties.title = draft.title
    path = Path(path)
    doc.save(path)
    return path
