"""Разметка документов для выгрузки (FR-7.1, FR-7.2, FR-8.3): DOCX, PDF и `Сводка.xlsx`.

На входе — находки из `db.exports.export_payload`: у каждой смещения `start`/`end` в тексте блока
и карта исходных позиций блока (`source_map`, контракт 2). По карте находка переводится в места
исходного файла:

- DOCX — отрезки runs (`DocxRunSpan`): run режется по границам находки, отрезок получает заливку
  цвета уровня, на весь диапазон ставится комментарий Word с текстом выноски (python-docx ≥ 1.2).
  Комментарии Word есть только в основной части документа: находки в сносках получают одну заливку.
- PDF — рамки слов (`WordBox`): слова одной строки объединяются в прямоугольник, на них ложится
  аннотация-выделение цвета уровня с текстом выноски во всплывающей заметке (PyMuPDF).

Системы координат PDF (проверено на PyMuPDF 2.0): `page.get_text` отдаёт рамки текстового слоя
в системе страницы без поворота — в ней же их принимает `add_highlight_annot`. Рамки слов OCR
считаются по изображению страницы, то есть в системе `page.rect` уже повёрнутой страницы; на странице
с `/Rotate` они переводятся обратно через `page.derotation_matrix`.

Находка, у которой нет ни одного отрезка в карте (текст из `w:sym`, пустая карта), не размещается:
она остаётся в сводке, а в журнал пишется предупреждение.
"""

from __future__ import annotations

import copy
import logging
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

log = logging.getLogger(__name__)

LEVEL_LABELS = {"critical": "Критично", "warning": "Внимание"}
SEVERITY = {"warning": 1, "critical": 2}
STATUS_LABELS = {"new": "Новая", "accepted": "Принята", "edited": "Изменена", "rejected": "Отклонена"}
AI_AUTHOR = "Светофор рисков (ИИ)"
AI_INITIALS = "ИИ"

# Заливка DOCX и цвет выделения PDF: светлые тона красного и жёлтого, текст под ними читается
DOCX_FILL = {"critical": "FFC7CE", "warning": "FFEB9C"}
PDF_COLOR = {"critical": (0.96, 0.42, 0.42), "warning": (1.0, 0.84, 0.18)}
XLSX_FILL = {"Критично": "FFC7CE", "Внимание": "FFEB9C", "Без рисков": "C6EFCE"}


@dataclass
class MarkResult:
    """Сколько находок легло в файл и какие не нашли места в исходнике."""

    placed: int = 0
    skipped: list[str] = field(default_factory=list)


def _worst(levels: Iterable[str]) -> str:
    return max(levels, key=lambda lv: SEVERITY.get(lv, 0))


def callout_lines(f: dict) -> list[str]:
    """Текст выноски: уровень и правило, пояснение, рекомендация, решение и комментарий юриста."""
    rule = f.get("rule")
    head = f"{LEVEL_LABELS.get(f['level'], f['level'])}: {rule['code']} {rule['title']}" if rule else \
        f"{LEVEL_LABELS.get(f['level'], f['level'])}: ручная пометка"
    if f.get("order"):
        head = f"{f['order']}. {head}"
    lines = [head]
    if f.get("explanation"):
        lines.append(f["explanation"])
    if f.get("recommendation"):
        lines.append(f"Рекомендация: {f['recommendation']}")
    if f.get("status") in ("accepted", "edited", "rejected"):
        who = (f.get("decided_by") or {}).get("full_name")
        lines.append(f"Статус: {STATUS_LABELS[f['status']]}" + (f" ({who})" if who else ""))
    if f.get("comment"):
        lines.append(f"Комментарий: {f['comment']}")
    return lines


def author_of(f: dict) -> tuple[str, str]:
    """Автор комментария и инициалы: находка ИИ — «Светофор рисков», ручная пометка — её автор."""
    if f.get("source") == "user" and f.get("author"):
        name = f["author"]["full_name"] or "Пользователь"
        initials = "".join(part[0] for part in name.split()[:2] if part).upper()
        return name, initials or "П"
    return AI_AUTHOR, AI_INITIALS


# --- DOCX -----------------------------------------------------------------------------------


def _w(tag: str) -> str:
    from docx.oxml.ns import qn

    return qn(f"w:{tag}")


def _content_length(child) -> int:
    """Сколько символов `Run.text` даёт элемент run — так же считает `ml.extract.docx_reader`."""
    tag = child.tag
    if tag == _w("t"):
        return len(child.text or "")
    if tag in (_w("tab"), _w("ptab"), _w("cr"), _w("noBreakHyphen")):
        return 1
    if tag == _w("br"):
        return 1 if child.get(_w("type")) in (None, "textWrapping") else 0
    return 0


def run_length(r) -> int:
    return sum(_content_length(c) for c in r if c.tag != _w("rPr"))


def split_run(r, cuts: Iterable[int]) -> list[tuple[int, int, Any]]:
    """Режет `w:r` по смещениям в его тексте; возвращает отрезки `(lo, hi, w:r)` по порядку.

    Каждый отрезок — копия run с тем же форматированием и своей частью содержимого. Элементы без
    текста (рисунок, поле, сноска) остаются в одном отрезке — в том, где они стоят.
    """
    total = run_length(r)
    bounds = [0, *sorted({c for c in cuts if 0 < c < total}), total]
    if len(bounds) == 2:
        return [(0, total, r)]
    pieces = []
    rpr, xml_space = _w("rPr"), "{http://www.w3.org/XML/1998/namespace}space"
    for lo, hi in zip(bounds, bounds[1:], strict=False):
        piece = copy.deepcopy(r)
        pos = 0
        for child in [c for c in piece if c.tag != rpr]:
            n = _content_length(child)
            start, end = pos, pos + n
            pos = end
            keep = (start < hi and end > lo) if n else (lo <= start < hi or (hi == total and start == total))
            if not keep:
                piece.remove(child)
            elif child.tag == _w("t") and (start < lo or end > hi):
                child.text = (child.text or "")[max(lo - start, 0):min(hi, end) - start]
                child.set(xml_space, "preserve")
        pieces.append((lo, hi, piece))
    parent = r.getparent()
    index = parent.index(r)
    parent.remove(r)
    for k, (_, _, piece) in enumerate(pieces):
        parent.insert(index + k, piece)
    return pieces


# Элементы `w:rPr`, которые по схеме идут после `w:shd`: заливку нужно вставить перед ними,
# иначе Word считает файл повреждённым
_RPR_AFTER_SHD = ("fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath",
                  "rPrChange")


def shade_run(r, fill: str) -> None:
    """Заливка run цветом уровня; своё выделение маркером у отрезка снимается, чтобы заливку было видно."""
    from docx.oxml import OxmlElement

    rpr = r.get_or_add_rPr()
    for old in [*rpr.findall(_w("shd")), *rpr.findall(_w("highlight"))]:
        rpr.remove(old)
    shd = OxmlElement("w:shd", {_w("val"): "clear", _w("color"): "auto", _w("fill"): fill})
    after = {_w(t) for t in _RPR_AFTER_SHD}
    for child in rpr:
        if child.tag in after:
            child.addprevious(shd)
            break
    else:
        rpr.append(shd)


@dataclass
class _Target:
    finding: int
    lo: int  # в тексте run
    hi: int


def mark_docx(src: str | Path, findings: list[dict], out: str | Path) -> MarkResult:
    """DOCX с заливкой рискованных фрагментов и комментариями Word (FR-7.1).

    DOC и RTF размечаются в DOCX, который сделал LibreOffice на этапе извлечения (`normalized_key`):
    карта runs построена по нему.
    """
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.parser import parse_xml
    from docx.text.run import Run
    from lxml import etree

    document = docx.Document(str(src))
    roots: dict[str, Any] = {"document": document.element}
    footnotes_part = None
    for rel in document.part.rels.values():
        if rel.reltype == RT.FOOTNOTES and not rel.is_external:
            footnotes_part = rel.target_part
            roots["footnotes"] = parse_xml(footnotes_part.blob)

    paragraphs: dict[str, list] = {}
    runs_of: dict[int, list] = {}

    def find_run(part: str, p_index: int, r_index: int):
        if part not in roots:
            return None
        if part not in paragraphs:
            paragraphs[part] = list(roots[part].iter(_w("p")))
        ps = paragraphs[part]
        if not 0 <= p_index < len(ps):
            return None
        p = ps[p_index]
        if id(p) not in runs_of:
            runs_of[id(p)] = list(p.iter(_w("r")))
        rs = runs_of[id(p)]
        return rs[r_index] if 0 <= r_index < len(rs) else None

    # Сначала все места всех находок: после разрезания run номера runs в абзаце меняются
    targets: dict[int, tuple[Any, list[_Target]]] = {}
    ranges: dict[int, tuple[tuple[Any, int], tuple[Any, int], str]] = {}  # находка → (первый run, lo), (последний, hi)
    result = MarkResult()
    for i, f in enumerate(findings):
        sm = f.get("source_map") or {}
        part = sm.get("part") or "document"
        first = last = None
        if sm.get("format") == "docx":
            for span in sm.get("runs") or []:
                a, b = max(span["start"], f["start"]), min(span["end"], f["end"])
                if a >= b:
                    continue
                r = find_run(part, span["paragraph_index"], span["run_index"])
                if r is None:
                    continue
                total = run_length(r)
                lo = span["run_offset"] + a - span["start"]
                hi = min(span["run_offset"] + b - span["start"], total)
                if lo >= hi:
                    continue
                targets.setdefault(id(r), (r, []))[1].append(_Target(i, lo, hi))
                first = first or (r, lo)
                last = (r, hi)
        if first is None:
            result.skipped.append(f["id"])
            log.warning("Выгрузка DOCX: находку %s не удалось разместить в исходнике (нет отрезков runs)", f["id"])
            continue
        ranges[i] = (first, last, part)

    # Разрезаем runs по границам находок и заливаем отрезки цветом худшего уровня
    pieces: dict[int, list[tuple[int, int, Any]]] = {}
    for key, (r, items) in targets.items():
        cuts = {t.lo for t in items} | {t.hi for t in items}
        split = split_run(r, cuts)
        pieces[key] = split
        for lo, hi, piece in split:
            covering = [findings[t.finding]["level"] for t in items if t.lo < hi and t.hi > lo]
            if covering:
                shade_run(piece, DOCX_FILL.get(_worst(covering), DOCX_FILL["warning"]))

    def piece_at(r, pos: int, *, end: bool):
        for lo, hi, piece in pieces[id(r)]:
            if (hi == pos) if end else (lo == pos):
                return piece
        return None

    for i, ((r0, lo), (r1, hi), part) in ranges.items():
        result.placed += 1
        if part != "document":
            continue  # комментарии Word python-docx ставит только в основной части
        first_piece, last_piece = piece_at(r0, lo, end=False), piece_at(r1, hi, end=True)
        if first_piece is None or last_piece is None:
            continue
        f = findings[i]
        lines = callout_lines(f)
        author, initials = author_of(f)
        comment = document.add_comment([Run(first_piece, None), Run(last_piece, None)], text=lines[0],
                                       author=author, initials=initials)
        for line in lines[1:]:
            comment.add_paragraph(line)

    if footnotes_part is not None:
        footnotes_part._blob = etree.tostring(roots["footnotes"], xml_declaration=True, encoding="UTF-8",
                                              standalone=True)
    document.save(str(out))
    return result


# --- PDF ------------------------------------------------------------------------------------


def _line_rects(words: list[dict]) -> list[tuple[float, float, float, float]]:
    """Слова подряд на одной строке → один прямоугольник: выделение выглядит как маркер, а не рябь."""
    rects: list[list[float]] = []
    for w in words:
        b = w["bbox"]
        x0, y0, x1, y1 = b["x0"], b["y0"], b["x1"], b["y1"]
        if rects:
            r = rects[-1]
            overlap = min(r[3], y1) - max(r[1], y0)
            if overlap > 0.5 * min(r[3] - r[1], y1 - y0) and x0 >= r[0] - 1:
                r[0], r[1], r[2], r[3] = min(r[0], x0), min(r[1], y0), max(r[2], x1), max(r[3], y1)
                continue
        rects.append([x0, y0, x1, y1])
    return [tuple(r) for r in rects]


def mark_pdf(src: str | Path, findings: list[dict], out: str | Path) -> MarkResult:
    """PDF с выделениями цвета уровня и всплывающими заметками; для сканов — по рамкам слов OCR (FR-7.2)."""
    import pymupdf

    result = MarkResult()
    with pymupdf.open(str(src)) as doc:
        for f in findings:
            sm = f.get("source_map") or {}
            words = sorted((w for w in sm.get("words") or [] if w["start"] < f["end"] and w["end"] > f["start"]),
                           key=lambda w: w["start"])
            by_page: dict[int, list[dict]] = {}
            for w in words:
                if 1 <= w["page"] <= doc.page_count:
                    by_page.setdefault(w["page"], []).append(w)
            if not by_page:
                result.skipped.append(f["id"])
                log.warning("Выгрузка PDF: находку %s не удалось разместить в исходнике (нет рамок слов)", f["id"])
                continue
            text = "\n".join(callout_lines(f))
            author, _ = author_of(f)
            color = PDF_COLOR.get(f["level"], PDF_COLOR["warning"])
            for page_no, page_words in by_page.items():
                page = doc[page_no - 1]
                rects = []
                for rect in _line_rects(page_words):
                    rect = pymupdf.Rect(rect)
                    # Рамки OCR — в системе повёрнутой страницы (`page.rect`), аннотации — без поворота
                    if page.rotation and page_words[0].get("confidence") is not None:
                        rect = rect * page.derotation_matrix
                    rects.append(rect)
                annot = page.add_highlight_annot(rects)
                annot.set_colors(stroke=color)
                annot.set_info(content=text, title=author, subject=LEVEL_LABELS.get(f["level"], ""))
                annot.update()
                anchor = rects[0]
                annot.set_popup(pymupdf.Rect(anchor.x1, anchor.y0, anchor.x1 + 220, anchor.y0 + 120))
                annot.set_open(False)
            result.placed += 1
        doc.save(str(out), garbage=3, deflate=True)
    return result


# --- Сводка.xlsx ----------------------------------------------------------------------------


def _local_time(value: str | None) -> str:
    """ISO-время из БД → «29.10.2026 14:05» по часовому поясу `EXPORT_TIMEZONE` (по умолчанию Москва)."""
    import os

    if not value:
        return ""
    moment = datetime.fromisoformat(value)
    try:
        from zoneinfo import ZoneInfo

        moment = moment.astimezone(ZoneInfo(os.environ.get("EXPORT_TIMEZONE") or "Europe/Moscow"))
        return moment.strftime("%d.%m.%Y %H:%M")
    except Exception:  # noqa: BLE001 — нет базы часовых поясов: время в UTC с пометкой
        return moment.strftime("%d.%m.%Y %H:%M UTC")


def summary_xlsx(summary: dict, out: str | Path, *, notes: Iterable[str] = ()) -> None:
    """`Сводка.xlsx` из `db.exports.batch_summary`: листы «Итоги», «Документы», «Находки» (FR-8.2, FR-8.3)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    header_fill = PatternFill("solid", fgColor="D9E1F2")

    def fill(cell, label: str | None) -> None:
        if label in XLSX_FILL:
            cell.fill = PatternFill("solid", fgColor=XLSX_FILL[label])

    def table(ws, headers: list[tuple[str, int]], rows: list[list], level_col: int | None = None) -> None:
        ws.append([h for h, _ in headers])
        for col, (_, width) in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font, cell.fill, cell.alignment = bold, header_fill, wrap
            ws.column_dimensions[cell.column_letter].width = width
        for row in rows:
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.alignment = wrap
            if level_col is not None:
                fill(ws.cell(row=ws.max_row, column=level_col), row[level_col - 1])
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions

    totals = summary["totals"]
    counters, review, found = totals.get("counters") or {}, totals.get("review") or {}, totals.get("findings") or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Итоги"
    items: list[tuple[str, Any]] = [
        ("Пакет", totals.get("name") or "Без названия"),
        ("Предмет закупки", totals.get("subject") or "—"),
        ("Профиль правил", totals.get("profile") or "—"),
        ("Автор", totals.get("author") or "—"),
        ("Создан", _local_time(totals.get("created_at"))),
        ("Итоговый уровень", totals.get("worst_level") or "—"),
        ("Документов", counters.get("total", 0)),
        ("Документов «Критично»", counters.get("critical", 0)),
        ("Документов «Внимание»", counters.get("warning", 0)),
        ("Документов «Без рисков»", counters.get("ok", 0)),
        ("Документов с ошибкой", counters.get("errors", 0)),
        ("Находок «Критично»", found.get("critical", 0)),
        ("Находок «Внимание»", found.get("warning", 0)),
        ("Проверено юристом", f"{review.get('checked', 0)} из {review.get('total', 0)}"),
    ]
    from db.exports import passport_fields

    passport = passport_fields(totals.get("meta"))  # {"fields": [...]}; список — пакеты до 29.10
    if passport:
        items.append(("", ""))
        items.append(("Паспорт закупки", ""))
        items += [(p.get("label") or p.get("key"), p.get("value")) for p in passport]
    total_security = ((totals.get("meta") or {}).get("finance") or {}).get("total_security")
    if total_security is not None:
        items.append(("Обеспечения всего, руб.", f"{total_security:,.2f}".replace(",", " ").replace(".", ",")))
    notes = list(notes)
    if notes:
        items.append(("", ""))
        items += [("Примечание", n) for n in notes]
    for label, value in items:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = bold
        ws.cell(row=ws.max_row, column=2).alignment = wrap
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        if row[0].value == "Итоговый уровень":
            fill(row[1], row[1].value)
        if row[0].value == "Паспорт закупки":
            row[0].fill = header_fill
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 80

    table(wb.create_sheet("Документы"), [
        ("Путь в архиве", 40), ("Файл", 32), ("Страниц", 10), ("Статус", 12), ("Уровень", 14),
        ("Критично", 10), ("Внимание", 10), ("Причина ошибки", 50),
    ], [[d["original_path"], d["filename"], d["pages"], d["status"], d["worst_level"], d["critical"], d["warning"],
         d["error"]] for d in summary["documents"]], level_col=5)

    def section(r: dict) -> str:
        parts = [f"п. {r['section_number']}" if r.get("section_number") else "", r.get("section_title") or ""]
        return " ".join(p for p in parts if p)

    table(wb.create_sheet("Находки"), [
        ("№", 6), ("Документ", 28), ("Раздел", 30), ("Страница", 10), ("Уровень", 12), ("Правило", 32),
        ("Цитата", 60), ("Пояснение", 60), ("Рекомендация", 50), ("Статус", 12), ("Решение принял", 22),
        ("Комментарий", 40),
    ], [[r["order"] or None, r["document"], section(r), r["page"], r["level"], r["rule"], r["quote"],
         r["explanation"], r["recommendation"], r["status"], r["decided_by"], r["comment"]]
        for r in summary["findings"]], level_col=5)
    wb.save(str(out))
