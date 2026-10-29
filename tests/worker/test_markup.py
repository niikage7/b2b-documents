"""Разметка выгрузки без БД: DOCX по карте runs, PDF по рамкам слов (в том числе /Rotate), Сводка.xlsx."""

from __future__ import annotations

import zipfile
from pathlib import Path

import docx
import pymupdf
import pytest
from docx.oxml.ns import qn
from openpyxl import load_workbook

import ml
from worker import markup
from worker.exports import archive_path, safe_filename

CLAUSE = "Поставщик обязан внести обеспечение исполнения контракта в размере 30 % от цены в течение 1 дня."


def finding(fid: str, block_text: str, quote: str, level: str, source_map: dict, **extra) -> dict:
    start = block_text.index(quote)
    return {"id": fid, "order": extra.pop("order", 1), "source": "ai", "status": "new", "level": level,
            "rule": {"code": "R-03", "title": "Размер обеспечения"}, "start": start, "end": start + len(quote),
            "quote": quote, "explanation": "Обеспечение выше 10 %.", "recommendation": "Запросить разъяснения.",
            "comment": None, "author": None, "decided_by": None, "source_map": source_map, **extra}


@pytest.fixture
def contract_docx(tmp_path) -> Path:
    """Пункт из нескольких runs с разным форматированием: границы находок режут runs посередине."""
    d = docx.Document()
    d.add_paragraph("Общие положения")
    p = d.add_paragraph()
    p.add_run("Поставщик обязан внести ")
    p.add_run("обеспечение исполнения контракта").bold = True
    p.add_run(" в размере 30 % от цены")
    p.add_run(" в течение 1 дня.").italic = True
    path = tmp_path / "contract.docx"
    d.save(path)
    return path


def paragraph(result, text: str):
    return next(p for p in result.paragraphs if p.text == text)


def shaded_text(path: Path) -> dict[str, str]:
    """Текст под заливкой по цветам, в порядке документа."""
    out: dict[str, str] = {}
    for r in docx.Document(str(path)).element.iter(qn("w:r")):
        shd = r.find(f"{qn('w:rPr')}/{qn('w:shd')}")
        if shd is not None:
            fill = shd.get(qn("w:fill"))
            out[fill] = out.get(fill, "") + "".join(t.text or "" for t in r.iter(qn("w:t")))
    return out


def test_docx_shading_and_comments_follow_run_map(contract_docx, tmp_path):
    extracted = ml.extract_docx(contract_docx)
    para = paragraph(extracted, CLAUSE)
    sm = para.source_map.model_dump()
    findings = [
        finding("f1", CLAUSE, "обеспечение исполнения контракта в размере 30 %", "critical", sm),
        # Перекрывается с первой: общий отрезок получает худший уровень
        finding("f2", CLAUSE, "30 % от цены в течение 1 дня", "warning", sm, order=2),
    ]
    out = tmp_path / "marked.docx"
    result = markup.mark_docx(contract_docx, findings, out)

    assert result.placed == 2 and not result.skipped
    shaded = shaded_text(out)
    assert shaded[markup.DOCX_FILL["critical"]] == "обеспечение исполнения контракта в размере 30 %"
    assert shaded[markup.DOCX_FILL["warning"]] == " от цены в течение 1 дня"
    # Текст и форматирование не изменились: повторное извлечение даёт тот же абзац
    again = ml.extract_docx(out)
    assert [p.text for p in again.paragraphs] == [p.text for p in extracted.paragraphs]
    marked = docx.Document(str(out))
    bold = "".join(r.text for r in marked.paragraphs[1].runs if r.bold)
    assert bold == "обеспечение исполнения контракта"

    comments = list(marked.comments)
    assert len(comments) == 2
    assert comments[0].author == markup.AI_AUTHOR
    assert comments[0].text.startswith("1. Критично: R-03 Размер обеспечения")
    assert "Рекомендация: Запросить разъяснения." in comments[0].text
    body = marked.element.xml
    assert body.count("commentRangeStart") == 2 and body.count("commentRangeEnd") == 2


def test_docx_user_finding_and_unplaceable(contract_docx, tmp_path):
    sm = paragraph(ml.extract_docx(contract_docx), CLAUSE).source_map.model_dump()
    user = finding("u1", CLAUSE, "в течение 1 дня", "warning", sm, source="user", status="accepted",
                   rule=None, explanation=None, recommendation=None, comment="Срок нереален",
                   author={"id": "x", "full_name": "Иванова Анна"}, decided_by={"id": "x", "full_name": "Иванова Анна"})
    lost = finding("lost", CLAUSE, "Поставщик", "critical", {"format": "docx", "runs": []})
    out = tmp_path / "marked.docx"
    result = markup.mark_docx(contract_docx, [user, lost], out)
    assert result.placed == 1 and result.skipped == ["lost"]
    comment = list(docx.Document(str(out)).comments)[0]
    assert comment.author == "Иванова Анна" and comment.initials == "ИА"
    assert "ручная пометка" in comment.text and "Комментарий: Срок нереален" in comment.text
    assert "Статус: Принята (Иванова Анна)" in comment.text


def test_split_run_keeps_tabs_and_breaks():
    r = docx.Document().add_paragraph().add_run("ab\tcd")._r
    pieces = markup.split_run(r, [1, 3])
    texts = [(lo, hi, "".join(t.text for t in p.iter(qn("w:t"))), len(p.findall(qn("w:tab"))))
             for lo, hi, p in pieces]
    assert texts == [(0, 1, "a", 0), (1, 3, "b", 1), (3, 5, "cd", 0)]


def _words_pdf(path: Path, *, rotation: int = 0) -> tuple[str, list[dict]]:
    """PDF с одной строкой и карта слов из `get_text("words")`, как у текстового слоя ML."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), "Penalty of five percent per day of delay", fontsize=12)
    if rotation:
        page.set_rotation(rotation)
    words, text, pos = [], "", 0
    for x0, y0, x1, y1, w, *_ in page.get_text("words"):
        if text:
            text += " "
            pos += 1
        words.append({"start": pos, "end": pos + len(w), "page": 1, "bbox": {"x0": x0, "y0": y0, "x1": x1, "y1": y1},
                      "confidence": None})
        text += w
        pos += len(w)
    doc.save(path)
    return text, words


def highlighted(path: Path) -> list[tuple[str, str, tuple]]:
    with pymupdf.open(path) as doc:
        page = doc[0]
        return [(a.type[1], page.get_textbox(a.rect).strip(), a.colors["stroke"], a.info["content"])
                for a in page.annots()]


@pytest.mark.parametrize("rotation", [0, 90])
def test_pdf_highlight_on_text_layer(tmp_path, rotation):
    src = tmp_path / "src.pdf"
    text, words = _words_pdf(src, rotation=rotation)
    f = finding("p1", text, "five percent per day", "critical", {"format": "pdf", "words": words})
    out = tmp_path / "out.pdf"
    assert markup.mark_pdf(src, [f], out).placed == 1
    [(kind, covered, color, content)] = highlighted(out)
    assert kind == "Highlight"
    assert "five percent per day" in covered and "Penalty" not in covered and "delay" not in covered
    assert color == pytest.approx(markup.PDF_COLOR["critical"], abs=0.01)
    assert content.startswith("1. Критично: R-03")


def test_pdf_ocr_words_on_rotated_page(tmp_path):
    """Рамки OCR — в системе повёрнутой страницы (`page.rect`): выделение всё равно ложится на слово."""
    src = tmp_path / "src.pdf"
    text, words = _words_pdf(src, rotation=90)
    with pymupdf.open(src) as doc:
        page = doc[0]
        for w in words:
            b = pymupdf.Rect(w["bbox"]["x0"], w["bbox"]["y0"], w["bbox"]["x1"], w["bbox"]["y1"]) * page.rotation_matrix
            w["bbox"] = {"x0": b.x0, "y0": b.y0, "x1": b.x1, "y1": b.y1}
            w["confidence"] = 0.93
    f = finding("o1", text, "percent", "warning", {"format": "pdf", "words": words})
    out = tmp_path / "out.pdf"
    markup.mark_pdf(src, [f], out)
    [(_, covered, color, _)] = highlighted(out)
    # Рамка аннотации чуть шире выделения: соседние буквы по краям попадают в get_textbox
    assert "percent" in covered and "five" not in covered and "day" not in covered
    assert color == pytest.approx(markup.PDF_COLOR["warning"], abs=0.01)


def test_summary_xlsx_sheets(tmp_path):
    summary = {
        "totals": {"name": "Закупка", "subject": "Поставка бумаги", "profile": "44-ФЗ", "author": "Тест",
                   "created_at": "2026-10-29T10:00:00+00:00", "worst_level": "Критично",
                   "meta": {"passport": {"fields": [{"key": "nmck", "label": "НМЦК", "value": "4 850 000,00 руб."}]}},
                   "counters": {"total": 2, "ok": 0, "warning": 1, "critical": 1, "errors": 0},
                   "review": {"checked": 1, "total": 3}, "findings": {"critical": 2, "warning": 1}},
        "documents": [{"filename": "Контракт.docx", "original_path": "docs/Контракт.docx", "pages": 3,
                       "status": "Готов", "worst_level": "Критично", "critical": 2, "warning": 0, "error": None}],
        "findings": [{"order": 1, "document": "Контракт.docx", "original_path": "docs/Контракт.docx",
                      "section_number": "3.1", "section_title": "Цена", "page": 2, "level": "Критично",
                      "rule": "R-01 Неустойка", "quote": "0,5 %", "explanation": "Высокая", "recommendation": None,
                      "status": "Принята", "decided_by": "Иванова", "comment": None}],
    }
    out = tmp_path / "Сводка.xlsx"
    markup.summary_xlsx(summary, out, notes=["docs/Приложение.pdf: не удалось разметить"])
    wb = load_workbook(out)
    assert wb.sheetnames == ["Итоги", "Документы", "Находки"]
    totals = {r[0]: r[1] for r in wb["Итоги"].iter_rows(values_only=True)}
    assert totals["Пакет"] == "Закупка" and totals["НМЦК"] == "4 850 000,00 руб."
    assert totals["Создан"] == "29.10.2026 13:00" or totals["Создан"].endswith("UTC")
    assert totals["Проверено юристом"] == "1 из 3"
    assert totals["Примечание"] == "docs/Приложение.pdf: не удалось разметить"
    rows = list(wb["Находки"].iter_rows(values_only=True))
    assert rows[0][:3] == ("№", "Документ", "Раздел")
    assert rows[1][2] == "п. 3.1 Цена" and rows[1][4] == "Критично"
    assert wb["Находки"]["E2"].fill.fgColor.rgb.endswith("FFC7CE")


def test_archive_path():
    used: set[str] = {"сводка.xlsx"}
    assert archive_path("docs/Договор.doc", "docx", used) == "docs/Договор.docx"
    assert archive_path("docs/Договор.docx", "docx", used) == "docs/Договор (2).docx"
    assert archive_path("../../etc/x.pdf", "pdf", used) == "etc/x.pdf"
    assert archive_path("Сводка.xlsx", None, used) == "Сводка (2).xlsx"


def test_marked_docx_is_valid_zip(contract_docx, tmp_path):
    """Части документа на месте: comments.xml добавлен, связи и типы содержимого записаны."""
    sm = paragraph(ml.extract_docx(contract_docx), CLAUSE).source_map.model_dump()
    out = tmp_path / "marked.docx"
    markup.mark_docx(contract_docx, [finding("f", CLAUSE, "30 %", "critical", sm)], out)
    with zipfile.ZipFile(out) as zf:
        names = set(zf.namelist())
        assert {"word/document.xml", "word/comments.xml", "[Content_Types].xml"} <= names
        assert "comments+xml" in zf.read("[Content_Types].xml").decode()


def test_safe_filename():
    assert safe_filename("Этап 7: выгрузка_разметка.zip") == "Этап 7_ выгрузка_разметка.zip"
    assert safe_filename(r'a/b\c*?"<>|.pdf') == "a_b_c_.pdf"
    assert safe_filename("..") == "выгрузка"
    assert safe_filename(None) is None
