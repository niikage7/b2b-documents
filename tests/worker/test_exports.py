"""Задачи выгрузки на настоящем PostgreSQL: размеченный DOCX и PDF, zip пакета со сводкой, ошибки.

Пакет проходит конвейер в режиме eager с подделкой LLM (`fake_llm`), затем выгрузка — теми же задачами,
что в очереди `export`. Сеть и настоящий LLM не используются.
"""

from __future__ import annotations

import io
import uuid
import zipfile

import docx
import pymupdf
from docx.oxml.ns import qn
from openpyxl import load_workbook

from bench import samples
from db import exports, findings
from db.models import Export, Finding
from tests.worker.conftest import make_batch, sample_bytes, zip_bytes
from tests.worker.test_analysis import ai_findings, contract, notice, risk_plan  # noqa: F401 — фикстуры
from tests.worker.test_pipeline import docs_of, run_batch
from worker import exports as export_tasks
from worker import markup


def export_events(bus, batch_id) -> list[dict]:
    return [data for _, type_, data in bus.events(str(batch_id)) if type_ == "export.done"]


def shaded(data: bytes) -> list[str]:
    out = []
    for r in docx.Document(io.BytesIO(data)).element.iter(qn("w:r")):
        if r.find(f"{qn('w:rPr')}/{qn('w:shd')}") is not None:
            out.append("".join(t.text or "" for t in r.iter(qn("w:t"))))
    return out


def test_document_export_docx_has_every_active_finding(s, env, author, fake_llm, bus, contract):  # noqa: F811
    fake_llm.plan = risk_plan("contract")
    batch_id = make_batch(s, env, author, [("Контракт.docx", contract)])
    run_batch(batch_id)
    doc = docs_of(s, batch_id)["Контракт.docx"]
    found = ai_findings(s, doc.id)
    assert len(found) >= 3
    # Отклонённая находка в выгрузку не попадает
    rejected = found[0]
    findings.update_finding(s, rejected.id, actor_id=author.id, version=rejected.version, status="rejected")

    job = exports.create_export(s, scope="document", target_id=doc.id, actor_id=author.id)
    assert job["format"] == "docx"
    view = export_tasks.export_document.apply(args=(job["export_id"],)).get()

    assert view["status"] == "done" and view["download_url"] == f"/api/exports/{job['export_id']}/download"
    assert view["filename"] == "Контракт_разметка.docx"
    data = env.objects[export_tasks.result_key(job["export_id"], "docx")]
    marked = docx.Document(io.BytesIO(data))
    assert len(list(marked.comments)) == len(found) - 1
    text = "".join(shaded(data))
    for f in found[1:]:
        assert f.quote in text
    assert export_events(bus, batch_id) == [view]


def test_document_export_pdf(s, env, author, fake_llm, fonts):
    pdf = sample_bytes(samples.make_contract_pdf, ".pdf")
    fake_llm.plan = risk_plan("contract")
    batch_id = make_batch(s, env, author, [("Контракт.pdf", pdf)])
    run_batch(batch_id)
    doc = docs_of(s, batch_id)["Контракт.pdf"]
    found = ai_findings(s, doc.id)
    assert found

    job = exports.create_export(s, scope="document", target_id=doc.id, actor_id=author.id)
    view = export_tasks.export_document.apply(args=(job["export_id"],)).get()
    assert view["status"] == "done" and view["filename"] == "Контракт_разметка.pdf"
    with pymupdf.open(stream=env.objects[export_tasks.result_key(job["export_id"], "pdf")], filetype="pdf") as out:
        annots = [(a.type[1], a.info["content"]) for page in out for a in page.annots()]
    assert len({content for _, content in annots}) == len(found)
    assert {kind for kind, _ in annots} == {"Highlight"}


def test_batch_export_zip_keeps_archive_structure_and_summary(s, env, author, fake_llm, contract, notice):  # noqa: F811
    fake_llm.plan = risk_plan("contract") | risk_plan("notice")
    archive = zip_bytes({"Закупка/Контракт.docx": contract, "Закупка/Приложения/Извещение.docx": notice,
                         "Закупка/битый.pdf": b"%PDF-1.4 broken"})
    batch_id = make_batch(s, env, author, [("пакет.zip", archive), ("Письмо.docx", notice)], name="Бумага")
    run_batch(batch_id)

    job = exports.create_export(s, scope="batch", target_id=batch_id, actor_id=author.id)
    assert job["format"] == "zip"
    view = export_tasks.export_batch.apply(args=(job["export_id"],)).get()
    assert view is None or view["status"] == "done"
    s.expire_all()
    export = s.get(Export, uuid.UUID(job["export_id"]))
    assert export.status == "done"
    assert exports.export_view(s, job["export_id"])["filename"] == "Бумага_разметка.zip"

    with zipfile.ZipFile(io.BytesIO(env.objects[export.object_key])) as zf:
        names = sorted(zf.namelist())
        # Битый файл с ошибкой в архив не попадает, он есть в сводке
        assert names == ["Закупка/Контракт.docx", "Закупка/Приложения/Извещение.docx", "Письмо.docx", "Сводка.xlsx"]
        assert shaded(zf.read("Закупка/Контракт.docx"))
        wb = load_workbook(io.BytesIO(zf.read("Сводка.xlsx")))
    assert wb.sheetnames == ["Итоги", "Документы", "Находки"]
    totals = {r[0]: r[1] for r in wb["Итоги"].iter_rows(values_only=True)}
    assert totals["Пакет"] == "Бумага" and totals["Документов"] == 4 and totals["Документов с ошибкой"] == 1
    assert "Примечание" not in totals
    findings_count = s.query(Finding).filter(Finding.batch_id == export.target_id).count()
    assert wb["Находки"].max_row - 1 == findings_count
    # Части документов удалены, остался только итоговый zip
    assert [k for k in env.objects if k.startswith(f"exports/{job['export_id']}/")] == [export.object_key]


def test_batch_export_puts_original_when_markup_fails(s, env, author, fake_llm, contract, monkeypatch):  # noqa: F811
    fake_llm.plan = risk_plan("contract")
    batch_id = make_batch(s, env, author, [("Контракт.docx", contract)], name="Сбой")
    run_batch(batch_id)

    def broken(*_args, **_kwargs):
        raise ValueError("сломанная разметка")

    monkeypatch.setattr(markup, "mark_docx", broken)
    job = exports.create_export(s, scope="batch", target_id=batch_id, actor_id=author.id)
    export_tasks.export_batch.apply(args=(job["export_id"],)).get()
    view = exports.export_view(s, job["export_id"])
    assert view["status"] == "done"
    with zipfile.ZipFile(io.BytesIO(env.objects[view["object_key"]])) as zf:
        assert zf.read("Контракт.docx") == contract
        totals = {r[0]: r[1] for r in load_workbook(io.BytesIO(zf.read("Сводка.xlsx")))["Итоги"].iter_rows(
            values_only=True)}
    assert "не удалось разметить" in totals["Примечание"]


def test_document_export_error_is_reported(s, env, author, fake_llm, bus, contract, monkeypatch):  # noqa: F811
    batch_id = make_batch(s, env, author, [("Контракт.docx", contract)])
    run_batch(batch_id)
    doc = docs_of(s, batch_id)["Контракт.docx"]

    def broken(*_args, **_kwargs):
        raise ValueError("сломанная разметка")

    monkeypatch.setattr(markup, "mark_docx", broken)
    job = exports.create_export(s, scope="document", target_id=doc.id, actor_id=author.id)
    assert export_tasks.export_document.apply(args=(job["export_id"],)).get() is None
    view = exports.export_view(s, job["export_id"])
    assert view["status"] == "error" and view["error_message"] == export_tasks.EXPORT_FAILED
    [event] = export_events(bus, batch_id)
    assert event["status"] == "error" and event["download_url"] is None
    assert any(r["task"] == "worker.exports.export_document" for r in bus.dead_letters())


def test_transient_error_is_retried(s, env, author, contract, monkeypatch):  # noqa: F811
    batch_id = make_batch(s, env, author, [("Контракт.docx", contract)])
    run_batch(batch_id)
    doc = docs_of(s, batch_id)["Контракт.docx"]
    real = markup.mark_docx
    calls = []

    def flaky(*args, **kwargs):
        calls.append(1)
        if len(calls) == 1:
            raise ConnectionError("MinIO недоступен")
        return real(*args, **kwargs)

    monkeypatch.setattr(markup, "mark_docx", flaky)
    monkeypatch.setattr(export_tasks, "retry_delay", lambda _n: 0)
    job = exports.create_export(s, scope="document", target_id=doc.id, actor_id=author.id)
    export_tasks.export_document.apply(args=(job["export_id"],))
    assert len(calls) == 2
    assert exports.export_view(s, job["export_id"])["status"] == "done"
