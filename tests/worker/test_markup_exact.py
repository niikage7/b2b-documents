"""Разметка выгрузки на стыке с парсером: карта runs и рамки слов из `ml` ложатся ровно на цитату.

`worker.markup` режет runs по `DocxRunSpan`, а их номера и смещения считает `ml.extract.docx_reader`.
Здесь случайные цитаты в блоках синтетических документов размечаются и файл читается тем же парсером
заново: залиты ровно символы цитат, текст не изменился. Для скана — подсветка по словам OCR.
"""

from __future__ import annotations

import random

import pymupdf
import pytest
from docx import Document

from bench.fonts import font_path
from bench.samples import make_contract_docx, make_contract_pdf, make_notice_docx
from bench.scangen import ScanStyle, make_scan
from ml.extract.docx_reader import extract_docx
from ml.extract.pdf_extract import extract_pdf
from ml.structure import build_structure
from worker.markup import DOCX_FILL, mark_docx, mark_pdf

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
needs_font = pytest.mark.skipif(font_path("serif") is None, reason="нет шрифта с кириллицей")


def finding(block, start: int, end: int, level: str = "warning", fid: str | None = None) -> dict:
    return {"id": fid or f"{block.order}:{start}", "order": None, "source": "ai", "status": "new", "level": level,
            "start": start, "end": end, "quote": block.text[start:end],
            "source_map": block.source_map.model_dump(mode="json"),
            "rule": {"code": "R-02", "title": "Обеспечение"}, "explanation": "Пояснение", "recommendation": None}


def fills_by_char(path) -> dict[int, dict[int, str]]:
    """Заливка уровня у каждого символа абзацев размеченного файла: {order: {позиция: цвет}}."""
    ours = set(DOCX_FILL.values())
    paragraphs = list(Document(str(path)).element.iter(f"{W}p"))
    runs: dict[int, list] = {}
    out: dict[int, dict[int, str]] = {}
    for p in extract_docx(path).paragraphs:
        chars = out.setdefault(p.order, {})
        for span in p.source_map.runs:
            r = runs.setdefault(span.paragraph_index, list(paragraphs[span.paragraph_index].iter(f"{W}r")))
            shd = r[span.run_index].find(f"{W}rPr/{W}shd")
            if shd is not None and shd.get(f"{W}fill") in ours:
                chars.update(dict.fromkeys(range(span.start, span.end), shd.get(f"{W}fill")))
    return out


@pytest.mark.parametrize("make", [make_contract_docx, make_notice_docx])
def test_docx_fill_matches_quotes_exactly(tmp_path, make):
    src = make(tmp_path / "src.docx")
    before = extract_docx(src)
    blocks = [b for b in build_structure(before, document_id="d") if b.source_map and len(b.text) > 10]
    rnd = random.Random(7)
    picked = rnd.sample(blocks, min(25, len(blocks)))
    findings = []
    for b in picked:
        start = rnd.randrange(0, len(b.text) - 3)
        findings.append(finding(b, start, rnd.randrange(start + 1, len(b.text) + 1)))
    # Пересечение в одном блоке: на общем отрезке — цвет худшего уровня
    first = findings[0]
    findings.append(finding(picked[0], first["start"], first["end"], level="critical", fid="overlap"))
    out = tmp_path / "out.docx"
    result = mark_docx(src, findings, out)
    assert result.placed == len(findings) and not result.skipped

    assert [p.text for p in extract_docx(out).paragraphs] == [p.text for p in before.paragraphs]
    fills = fills_by_char(out)
    for f, b in zip(findings, picked, strict=False):
        mapped = {i for s in b.source_map.runs for i in range(s.start, s.end)}
        want = set(range(f["start"], f["end"])) & mapped
        assert set(fills[b.order]) == want, f["quote"]
    assert set(fills[picked[0].order].values()) == {DOCX_FILL["critical"]}
    touched = {b.order for b in picked}
    assert not any(chars for order, chars in fills.items() if order not in touched)
    assert len(list(Document(str(out)).comments)) == len(findings)


@needs_font
@pytest.mark.tesseract
def test_scan_highlight_uses_ocr_words(tmp_path):
    contract = make_contract_pdf(tmp_path / "contract.pdf")
    scan = make_scan(contract, tmp_path / "scan.pdf", style=ScanStyle(max_skew=0.3), pages=[1])
    blocks = build_structure(extract_pdf(scan), document_id="d")
    block = max(blocks, key=lambda b: len(b.text))
    assert block.source_map.words and block.source_map.words[0].confidence is not None
    out = tmp_path / "out.pdf"
    assert mark_pdf(scan, [finding(block, 0, min(len(block.text), 30))], out).placed == 1
    with pymupdf.open(out) as doc:
        page = doc[block.page - 1]  # аннотация живёт, пока жива страница
        (annot,) = list(page.annots())
        first = block.source_map.words[0].bbox
        assert annot.rect.intersects(pymupdf.Rect(first.x0, first.y0, first.x1, first.y1))


@needs_font
@pytest.mark.tesseract
@pytest.mark.parametrize("rotation", [90, 270])
def test_real_ocr_words_on_rotated_page(tmp_path, rotation):
    """Настоящие рамки `ml.ocr_pages` на странице с /Rotate: подсветка ложится на слово текстового слоя."""
    from ml.ocr.pages import ocr_pages

    font = pymupdf.Font(fontfile=str(font_path("serif")))
    doc = pymupdf.open()
    page = doc.new_page()
    tw = pymupdf.TextWriter(page.rect)
    for i in range(6):
        tw.append((72, 100 + 30 * i), "Поставка в течение трёх календарных дней с даты заключения Контракта",
                  font=font, fontsize=14)
    tw.write_text(page)
    page.set_rotation(rotation)
    src = tmp_path / "rot.pdf"
    doc.save(src)
    (ocr,) = ocr_pages(src, [1], use_cache=False)
    word = next(w for w in ocr.words if w.text.startswith("Контракт"))
    text = word.text
    f = {"id": "w", "order": None, "source": "ai", "status": "new", "level": "critical", "start": 0,
         "end": len(text), "rule": None, "explanation": "", "source_map": {"format": "pdf", "words": [
             {"start": 0, "end": len(text), "page": 1, "bbox": word.bbox.model_dump(), "confidence": word.confidence}]}}
    out = tmp_path / "out.pdf"
    assert mark_pdf(src, [f], out).placed == 1
    with pymupdf.open(out) as marked:
        page = marked[0]
        (annot,) = list(page.annots())
        assert "Контракт" in page.get_textbox(annot.rect)
