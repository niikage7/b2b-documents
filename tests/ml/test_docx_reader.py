from __future__ import annotations

import docx
import pytest
from docx.oxml.ns import qn

from bench.docx_xml import lvl, make_docx, numbering, numpr, p, r
from ml.extract.docx_reader import extract_docx

MULTI = (
    lvl(0, "decimal", "%1.")
    + lvl(1, "decimal", "%1.%2.")
    + lvl(2, "decimal", "%1.%2.%3.")
)


def texts(result):
    return [(item.number, item.text) for item in result.paragraphs]


def test_multilevel_numbering_and_restart(tmp_path):
    body = "".join(
        [
            p("Общие положения", ppr=numpr(1, 0)),
            p("Термины", ppr=numpr(1, 1)),
            p("Подпункт", ppr=numpr(1, 2)),
            p("Подпункт два", ppr=numpr(1, 2)),
            p("Ещё пункт", ppr=numpr(1, 1)),
            p("Предмет контракта", ppr=numpr(1, 0)),
            p("Первый пункт раздела", ppr=numpr(1, 1)),
        ]
    )
    path = make_docx(tmp_path / "a.docx", body, numbering_xml=numbering({0: MULTI}, {1: 0}))
    assert texts(extract_docx(path)) == [
        ("1.", "Общие положения"),
        ("1.1.", "Термины"),
        ("1.1.1.", "Подпункт"),
        ("1.1.2.", "Подпункт два"),
        ("1.2.", "Ещё пункт"),
        ("2.", "Предмет контракта"),
        ("2.1.", "Первый пункт раздела"),
    ]


def test_shared_abstract_continues_and_start_override_restarts(tmp_path):
    override = '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>'
    body = "".join(
        [
            p("а", ppr=numpr(1)),
            p("б", ppr=numpr(1)),
            p("Текст между списками"),
            p("в — продолжение через другой num", ppr=numpr(2)),
            p("заново", ppr=numpr(3)),
            p("дальше", ppr=numpr(3)),
        ]
    )
    xml = numbering({0: lvl(0, "decimal", "%1)")}, {1: 0, 2: 0, 3: (0, override)})
    result = extract_docx(make_docx(tmp_path / "b.docx", body, numbering_xml=xml))
    assert [i.number for i in result.paragraphs] == ["1)", "2)", None, "3)", "1)", "2)"]


def test_formats_russian_roman_bullet_and_numid_zero(tmp_path):
    levels = {
        0: lvl(0, "russianLower", "%1)"),
        1: lvl(0, "upperRoman", "Раздел %1."),
        2: lvl(0, "bullet", "•"),
    }
    body = "".join(
        [p(f"буква {i}", ppr=numpr(1)) for i in range(7)]
        + [p("римский", ppr=numpr(2)), p("римский 2", ppr=numpr(2))]
        + [p("маркер", ppr=numpr(3))]
        + [p("без номера", ppr=numpr(0))]
    )
    xml = numbering(levels, {1: 0, 2: 1, 3: 2})
    result = extract_docx(make_docx(tmp_path / "c.docx", body, numbering_xml=xml))
    numbers = [i.number for i in result.paragraphs]
    # а б в г д е ж — без ё; й пропускается дальше в алфавите
    assert numbers[:7] == ["а)", "б)", "в)", "г)", "д)", "е)", "ж)"]
    assert numbers[7:9] == ["Раздел I.", "Раздел II."]
    assert numbers[9] is None and result.paragraphs[9].is_bullet
    assert numbers[10] is None and not result.paragraphs[10].is_bullet


def test_russian_letters_skip_like_word():
    from ml.extract.numbering import format_number

    assert [format_number(n, "russianLower") for n in (9, 10, 28, 29)] == ["и", "к", "я", "аа"]


def test_numbering_from_style_and_is_lgl(tmp_path):
    styles = (
        '<w:style w:type="paragraph" w:styleId="Clause"><w:name w:val="Пункт"/>'
        '<w:pPr><w:numPr><w:ilvl w:val="1"/><w:numId w:val="5"/></w:numPr></w:pPr></w:style>'
    )
    levels = lvl(0, "upperRoman", "%1.") + lvl(1, "decimal", "%1.%2.", extra="<w:isLgl/>")
    body = "".join(
        [
            p("Раздел", ppr=numpr(5, 0)),
            p("Пункт по стилю", ppr='<w:pStyle w:val="Clause"/>'),
            p("Ещё пункт по стилю", ppr='<w:pStyle w:val="Clause"/>'),
        ]
    )
    xml = numbering({0: levels}, {5: 0})
    result = extract_docx(make_docx(tmp_path / "d.docx", body, numbering_xml=xml, styles_extra=styles))
    assert texts(result) == [("I.", "Раздел"), ("1.1.", "Пункт по стилю"), ("1.2.", "Ещё пункт по стилю")]


def test_lvl_restart_zero_keeps_counting(tmp_path):
    levels = lvl(0, "decimal", "Статья %1.") + lvl(1, "decimal", "%2)", extra='<w:lvlRestart w:val="0"/>')
    body = "".join(
        [
            p("A", ppr=numpr(1, 0)),
            p("a1", ppr=numpr(1, 1)),
            p("B", ppr=numpr(1, 0)),
            p("b1", ppr=numpr(1, 1)),
        ]
    )
    result = extract_docx(make_docx(tmp_path / "e.docx", body, numbering_xml=numbering({0: levels}, {1: 0})))
    assert [i.number for i in result.paragraphs] == ["Статья 1.", "1)", "Статья 2.", "2)"]


def test_num_style_link(tmp_path):
    styles = (
        '<w:style w:type="numbering" w:styleId="ContractList"><w:name w:val="Contract List"/>'
        '<w:pPr><w:numPr><w:numId w:val="7"/></w:numPr></w:pPr></w:style>'
    )
    xml = numbering(
        {
            0: lvl(0, "decimal", "%1.") + '<w:styleLink w:val="ContractList"/>',
            1: '<w:numStyleLink w:val="ContractList"/>',
        },
        {7: 0, 8: 1},
    )
    body = p("первый", ppr=numpr(8)) + p("второй", ppr=numpr(8))
    result = extract_docx(make_docx(tmp_path / "f.docx", body, numbering_xml=xml, styles_extra=styles))
    assert [i.number for i in result.paragraphs] == ["1.", "2."]


def test_headings_by_outline_and_style(tmp_path):
    body = "".join(
        [
            p("Заголовок по стилю «1»", ppr='<w:pStyle w:val="1"/>'),
            p("Второй уровень", ppr='<w:pStyle w:val="Heading2"/>'),
            p("Обычный текст пункта, длинный и не жирный."),
            p("ЗАГОЛОВОК ПРОПИСНЫМИ"),
        ]
    )
    result = extract_docx(make_docx(tmp_path / "g.docx", body))
    heads = [(i.heading_level, i.heading_source) for i in result.paragraphs]
    assert heads == [(1, "outline"), (2, "style"), (None, None), (1, "heuristic")]
    assert result.paragraphs[0].bold


def test_tracked_changes_fields_links_sdt(tmp_path):
    runs = (
        r("Срок поставки ")
        + '<w:del w:id="1" w:author="x"><w:r><w:delText>30</w:delText></w:r></w:del>'
        + '<w:ins w:id="2" w:author="x"><w:r><w:t>3</w:t></w:r></w:ins>'
        + r(" дня, см. ")
        + '<w:hyperlink r:id="rId9"><w:r><w:t>приложение</w:t></w:r></w:hyperlink>'
        + '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        + '<w:r><w:instrText> PAGE </w:instrText></w:r>'
        + '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        + r(" 5")
        + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )
    sdt = '<w:sdt><w:sdtContent><w:p><w:r><w:t>Текст в элементе управления</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    result = extract_docx(make_docx(tmp_path / "h.docx", p(runs=runs) + sdt))
    assert [i.text for i in result.paragraphs] == [
        "Срок поставки 3 дня, см. приложение 5",
        "Текст в элементе управления",
    ]


def test_run_map_matches_python_docx(tmp_path):
    runs = (
        r("  Поставка ")
        + r("Товара", "<w:b/>")
        + '<w:r><w:tab/><w:t>в течение</w:t><w:br/><w:t>3 дней</w:t></w:r>'
        + '<w:ins w:id="1" w:author="x"><w:r><w:t xml:space="preserve"> с даты</w:t></w:r></w:ins>'
        + '<w:del w:id="2" w:author="x"><w:r><w:delText>удалено</w:delText></w:r></w:del>'
        + r(" заключения  ")
    )
    path = make_docx(tmp_path / "i.docx", p("Первый") + p(runs=runs))
    result = extract_docx(path)
    item = result.paragraphs[1]
    assert item.text == "Поставка Товара\tв течение\n3 дней с даты заключения"

    document = docx.Document(str(path))
    paragraphs = list(document.element.body.iter(qn("w:p")))
    covered = 0
    for span in item.source_map.runs:
        para = paragraphs[span.paragraph_index]
        run = list(para.iter(qn("w:r")))[span.run_index]
        run_text = docx.text.run.Run(run, None).text
        length = span.end - span.start
        assert item.text[span.start : span.end] == run_text[span.run_offset : span.run_offset + length]
        covered += length
    assert covered == len(item.text)


def test_tables_merges_and_layout(tmp_path):
    def tc(text, tcpr=""):
        return f"<w:tc>{f'<w:tcPr>{tcpr}</w:tcPr>' if tcpr else ''}{p(text)}</w:tc>"

    grid_table = (
        '<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/></w:tblPr>'
        f"<w:tr>{tc('Наименование', '<w:gridSpan w:val=\"2\"/>')}{tc('Срок', '<w:vMerge w:val=\"restart\"/>')}</w:tr>"
        f"<w:tr>{tc('Сервер')}{tc('2 шт.')}{tc('', '<w:vMerge/>')}</w:tr>"
        "</w:tbl>"
    )
    layout_table = f"<w:tbl><w:tr>{tc('Заказчик:')}{tc('Поставщик:')}</w:tr></w:tbl>"
    result = extract_docx(make_docx(tmp_path / "j.docx", grid_table + layout_table))
    cells = [(i.text, i.table.table_id, i.table.row, i.table.col, i.table.row_span, i.table.col_span, i.table.is_layout)
             for i in result.paragraphs]
    assert cells == [
        ("Наименование", 0, 0, 0, 1, 2, False),
        ("Срок", 0, 0, 2, 2, 1, False),
        ("Сервер", 0, 1, 0, 1, 1, False),
        ("2 шт.", 0, 1, 1, 1, 1, False),
        ("Заказчик:", 1, 0, 0, 1, 1, True),
        ("Поставщик:", 1, 0, 1, 1, 1, True),
    ]


def test_label_value_row_is_one_block(tmp_path):
    def tc(text, rpr=""):
        return f"<w:tc>{p(runs=r(text, rpr))}</w:tc>"

    bold = "<w:b/>"
    table = (
        "<w:tbl>"
        f"<w:tr>{tc('Заказчик', bold)}{tc('Поставщик', bold)}</w:tr>"
        f"<w:tr>{tc('Размер обеспечения заявки')}{tc('7789.41 РОССИЙСКИЙ РУБЛЬ')}<w:tc>{p()}</w:tc></w:tr>"
        f"<w:tr>{tc('Цена:')}{tc('1 000 руб.')}</w:tr>"
        f"<w:tr>{tc('От имени Заказчика:')}{tc('От имени Поставщика:')}</w:tr>"
        f"<w:tr>{tc('Подпись')}{tc('________ / Иванов /')}</w:tr>"
        "</w:tbl>"
    )
    path = make_docx(tmp_path / "lv.docx", table)
    result = extract_docx(path)
    assert [(i.text, i.table.row, i.table.col_span) for i in result.paragraphs] == [
        ("Заказчик", 0, 1),
        ("Поставщик", 0, 1),
        ("Размер обеспечения заявки: 7789.41 РОССИЙСКИЙ РУБЛЬ", 1, 2),
        ("Цена: 1 000 руб.", 2, 2),
        ("От имени Заказчика:", 3, 1),
        ("От имени Поставщика:", 3, 1),
        ("Подпись", 4, 1),
        ("________ / Иванов /", 4, 1),
    ]
    # Отрезки runs указывают на подпись и на значение в своих ячейках
    item = result.paragraphs[2]
    paragraphs = list(docx.Document(str(path)).element.body.iter(qn("w:p")))
    for span in item.source_map.runs:
        run = list(paragraphs[span.paragraph_index].iter(qn("w:r")))[span.run_index]
        run_text = docx.text.run.Run(run, None).text
        assert item.text[span.start : span.end] == run_text[span.run_offset : span.run_offset + span.end - span.start]
    assert sum(s.end - s.start for s in item.source_map.runs) == len(item.text) - 2  # кроме «: »


def test_typed_clause_numbers(tmp_path):
    body = (
        p("Ответственность Сторон", ppr="")
        + p(runs=r("6.4. ") + r("За каждый факт неисполнения"))
        + p("а) по заявке Заказчика;")
        + p("1.5 млн руб.")
        + "<w:tbl><w:tr><w:tc>" + p("1. Стол") + "</w:tc></w:tr></w:tbl>"
    )
    path = make_docx(tmp_path / "num.docx", body)
    result = extract_docx(path)
    assert [(i.number, i.number_source, i.text) for i in result.paragraphs] == [
        (None, None, "Ответственность Сторон"),
        ("6.4.", "pattern", "За каждый факт неисполнения"),
        ("а)", "pattern", "по заявке Заказчика;"),
        ("1.5", "pattern", "млн руб."),  # тот же шаблон, что в PDF: десятичное число в начале абзаца — номер
        (None, None, "1. Стол"),  # в ячейках таблиц номер не выделяется
    ]
    (span,) = result.paragraphs[1].source_map.runs
    assert (span.start, span.end, span.run_index, span.run_offset) == (0, 27, 1, 0)


def test_cell_with_several_paragraphs(tmp_path):
    cell = f"<w:tc>{p('Первый абзац')}{p('Второй', ppr=numpr(1))}</w:tc>"
    xml = numbering({0: lvl(0, "decimal", "%1.")}, {1: 0})
    result = extract_docx(make_docx(tmp_path / "k.docx", f"<w:tbl><w:tr>{cell}</w:tr></w:tbl>", numbering_xml=xml))
    (item,) = result.paragraphs
    assert item.text == "Первый абзац\n1. Второй"
    assert {s.paragraph_index for s in item.source_map.runs} == {0, 1}


def test_textbox_and_formula_are_separate(tmp_path):
    textbox = (
        '<w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData>'
        "<wps:wsp><wps:txbx><w:txbxContent>"
        + p("Текст надписи")
        + "</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        '<mc:Fallback><w:pict><w:txbxContent>' + p("Текст надписи") + "</w:txbxContent></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r>"
    )
    formula = "<m:oMath><m:r><m:t>x+1</m:t></m:r></m:oMath>"
    result = extract_docx(make_docx(tmp_path / "l.docx", p(runs=r("Основной текст") + textbox + formula)))
    assert [(i.stream, i.text) for i in result.paragraphs] == [
        ("body", "Основной текст"),
        ("textbox", "Текст надписи"),
        ("formula", "x+1"),
    ]


def test_pages_from_rendered_breaks(tmp_path):
    body = (
        p("Стр 1")
        + p(runs='<w:r><w:lastRenderedPageBreak/><w:t>Стр 2</w:t></w:r>')
        + p(runs='<w:r><w:t>Начало</w:t><w:lastRenderedPageBreak/><w:t> конец</w:t></w:r>')
    )
    result = extract_docx(make_docx(tmp_path / "m.docx", body))
    assert [(i.page, i.page_end) for i in result.paragraphs] == [(1, None), (2, None), (2, 3)]
    assert result.pages == 3 and not result.pages_estimated


def test_table_row_across_pages_counts_once(tmp_path):
    mark = "<w:lastRenderedPageBreak/>"

    def cell(*runs):
        return "<w:tc>" + p(runs="".join(runs)) + "</w:tc>"

    long_text = "описание " * 80
    table = (
        "<w:tbl>"
        f"<w:tr>{cell(r('1'))}{cell(r('Стол'))}{cell(r(long_text))}</w:tr>"
        # Строка переходит на страницу 2: метка в каждой ячейке и две в одной ячейке
        f"<w:tr>{cell(r('2'))}{cell(r('Шкаф'), f'<w:r>{mark}<w:t>для документов</w:t></w:r>')}"
        f"{cell(f'<w:r>{mark}</w:r>', r('Корпус'), f'<w:r>{mark}<w:t>ЛДСП</w:t></w:r>')}</w:tr>"
        # Повтор метки в начале следующей строки
        f"<w:tr>{cell(f'<w:r>{mark}<w:t>3</w:t></w:r>')}{cell(r('Тумба'))}{cell(r('ЛДСП'))}</w:tr>"
        "</w:tbl>"
    )
    result = extract_docx(make_docx(tmp_path / "rows.docx", p("До таблицы") + table + p("После таблицы")))
    pages = [(i.text[:12], i.page, i.page_end) for i in result.paragraphs]
    assert pages == [
        ("До таблицы", 1, None), ("1", 1, None), ("Стол", 1, None), (long_text[:12], 1, None),
        ("2", 1, None), ("Шкафдля доку", 1, 2), ("КорпусЛДСП", 2, None),  # вторая метка в ячейке — повтор
        ("3", 2, None), ("Тумба", 2, None), ("ЛДСП", 2, None),
        ("После таблиц", 2, None),
    ]
    assert result.pages == 2


def test_pages_calibrated_to_source_page_count(tmp_path):
    body = "".join(p("Пункт договора " * 40) for _ in range(30))
    path = make_docx(tmp_path / "doc.docx", body)
    assert extract_docx(path).pages > 5
    result = extract_docx(path, kind="doc", source_pages=3)
    assert result.pages == 3 and max(i.page for i in result.paragraphs) == 3


def test_pages_estimated_with_explicit_break(tmp_path):
    body = p("Стр 1") + p(runs='<w:r><w:br w:type="page"/></w:r>') + p("Стр 2")
    result = extract_docx(make_docx(tmp_path / "n.docx", body))
    assert [i.page for i in result.paragraphs] == [1, 2]
    assert result.pages_estimated and all(i.page_estimated for i in result.paragraphs)


def test_footnotes_and_non_bmp(tmp_path):
    footnotes = (
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:type="separator" w:id="-1"><w:p/></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t> Сноска о ГОСТ</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    )
    body = p(runs=r("Текст 😀") + '<w:r><w:footnoteReference w:id="1"/></w:r>')
    result = extract_docx(make_docx(tmp_path / "o.docx", body, footnotes_xml=footnotes))
    assert [(i.stream, i.text) for i in result.paragraphs] == [("body", "Текст �"), ("footnote", "Сноска о ГОСТ")]
    assert result.paragraphs[1].source_map.part == "footnotes"


@pytest.mark.parametrize("bad", [b"", b"PK\x03\x04garbage"])
def test_corrupted(tmp_path, bad):
    from ml.errors import CorruptedFileError

    path = tmp_path / "bad.docx"
    path.write_bytes(bad)
    with pytest.raises(CorruptedFileError):
        extract_docx(path)
