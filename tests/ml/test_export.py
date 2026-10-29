"""Разметка исходника находками (FR-7.1, FR-7.2): подсветка ложится ровно на цитату, текст не меняется."""

from __future__ import annotations

import random

import pymupdf
import pytest
from docx import Document

from bench.fonts import font_path
from bench.samples import make_contract_docx, make_contract_pdf, make_notice_docx
from bench.scangen import ScanStyle, make_scan
from ml.export import DOCX_FILL, annotate, annotate_docx, annotate_pdf, finding_note
from ml.extract.docx_reader import extract_docx
from ml.extract.pdf_extract import extract_pdf
from ml.structure import build_structure

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
needs_font = pytest.mark.skipif(font_path("serif") is None, reason="нет шрифта с кириллицей")


def finding(block, start, end, level="warning", **extra):
    return {"id": f"{block.order}:{start}", "level": level, "status": "new", "start": start, "end": end,
            "quote": block.text[start:end], "source_map": block.source_map.model_dump(), "bbox": block.bbox,
            "page": block.page, "section_number": block.section_number, "section_title": block.section_title,
            "rule": {"code": "R-02", "title": "Обеспечение исполнения контракта"},
            "explanation": "Размер обеспечения выше порога", "recommendation": "Запросить разъяснения"} | extra


def fills_by_char(path) -> dict[int, dict[int, str]]:
    """Заливка каждого символа абзацев: {order: {позиция: fill}} — через повторное извлечение файла."""
    result = extract_docx(path)
    paragraphs = list(Document(str(path)).element.iter(f"{W}p"))
    out: dict[int, dict[int, str]] = {}
    for p in result.paragraphs:
        chars = out.setdefault(p.order, {})
        for span in p.source_map.runs:
            run = list(paragraphs[span.paragraph_index].iter(f"{W}r"))[span.run_index]
            shd = run.find(f"{W}rPr/{W}shd")
            if shd is not None:
                chars.update(dict.fromkeys(range(span.start, span.end), shd.get(f"{W}fill")))
    return out


def comments(path) -> list[str]:
    doc = Document(str(path))
    return ["\n".join(p.text for p in c.paragraphs) for c in doc.comments]


@pytest.mark.parametrize("make", [make_contract_docx, make_notice_docx])
def test_docx_highlight_is_exact_and_text_unchanged(tmp_path, make):
    src = make(tmp_path / "src.docx")
    before = extract_docx(src)
    blocks = [b for b in build_structure(before, document_id="d") if b.source_map and len(b.text) > 10]
    rnd = random.Random(7)
    picked = rnd.sample(blocks, min(25, len(blocks)))
    findings = []
    for b in picked:
        start = rnd.randrange(0, len(b.text) - 3)
        findings.append(finding(b, start, rnd.randrange(start + 1, len(b.text) + 1)))
    out = tmp_path / "out.docx"
    stats = annotate_docx(src, findings, out)
    assert (stats.placed, stats.comments, stats.missed) == (len(findings), len(findings), [])

    after = extract_docx(out)
    assert [p.text for p in after.paragraphs] == [p.text for p in before.paragraphs]
    fills = fills_by_char(out)
    touched = {f["id"].split(":")[0] for f in findings}
    for f, b in zip(findings, picked, strict=True):
        mapped = {i for s in b.source_map.runs for i in range(s.start, s.end)}
        assert set(fills[b.order]) == set(range(f["start"], f["end"])) & mapped, f["quote"]
    assert not any(chars for order, chars in fills.items() if str(order) not in touched)
    notes = comments(out)
    assert len(notes) == len(findings)
    assert notes[0].startswith("Внимание · R-02 Обеспечение исполнения контракта")
    assert "Рекомендация: Запросить разъяснения" in notes[0]


def test_docx_cuts_runs_inside_words_tabs_and_breaks(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Поставщик обязан")
    p.add_run(" уплатить\tштраф").bold = True
    run = p.add_run(" 10 ")
    run.add_break()
    p.add_run("процентов цены")
    src = tmp_path / "src.docx"
    doc.save(src)
    (block,) = build_structure(extract_docx(src), document_id="d")
    text = block.text
    assert text == "Поставщик обязан уплатить\tштраф 10 \nпроцентов цены"
    warning = finding(block, text.index("обязан") + 2, text.index("штраф") + 3)
    critical = finding(block, text.index("штраф"), text.index("процентов") + 4, level="critical")
    out = tmp_path / "out.docx"
    annotate_docx(src, [critical, warning], out)
    fills = fills_by_char(out)[block.order]
    assert {i for i, f in fills.items() if f == DOCX_FILL["warning"]} == set(range(warning["start"], critical["start"]))
    # На пересечении — худший уровень, заливка критичной находки кладётся последней
    assert {i for i, f in fills.items() if f == DOCX_FILL["critical"]} == set(range(critical["start"], critical["end"]))
    assert extract_docx(out).paragraphs[0].text == text
    bold = [r for r in Document(str(out)).paragraphs[0].runs if r.bold]
    assert "".join(r.text for r in bold) == " уплатить\tштраф"  # формат кусков разрезанного run сохранён


def test_note_for_user_mark_and_decision():
    lines = finding_note({"level": "critical", "rule": None, "author": {"full_name": "Юрист Ю."},
                          "explanation": "Срок оплаты", "comment": "Срок оплаты", "status": "accepted",
                          "decided_by": {"full_name": "Юрист Ю."}, "section_number": "4.2", "page": 3})
    assert lines == ["Критично · ручная пометка (Юрист Ю.)", "Раздел: п. 4.2, стр. 3", "Срок оплаты",
                     "Статус: Принята (Юрист Ю.)"]


def test_finding_without_map_is_reported(tmp_path):
    src = make_contract_docx(tmp_path / "src.docx")
    stats = annotate_docx(src, [{"id": "x", "level": "warning", "start": 0, "end": 5, "source_map": None}],
                          tmp_path / "out.docx")
    assert (stats.placed, stats.missed) == (0, ["x"])


@needs_font
def test_pdf_highlight_covers_quote(tmp_path):
    src = make_contract_pdf(tmp_path / "contract.pdf")
    blocks = build_structure(extract_pdf(src, run_ocr=False), document_id="d")
    block = max(blocks, key=lambda b: len(b.text))
    f = finding(block, 10, 60, level="critical")
    out = tmp_path / "out.pdf"
    stats = annotate(src, [f], out)
    assert stats.placed == 1 and not stats.missed
    with pymupdf.open(out) as doc:
        # Аннотация живёт, пока жива её страница: страницы держатся в переменных
        assert [p.number + 1 for p in doc if p.first_annot] == [block.page]
        page = doc[block.page - 1]
        annots = list(page.annots())
        under = " ".join(page.get_textbox(a.rect) for a in annots).replace("\n", " ")
        assert f["quote"].split()[1] in under
        info = annots[0].info
        assert info["title"] == "Светофор рисков" and "Размер обеспечения выше порога" in info["content"]


@needs_font
def test_pdf_falls_back_to_block_bbox(tmp_path):
    src = make_contract_pdf(tmp_path / "contract.pdf")
    block = next(b for b in build_structure(extract_pdf(src, run_ocr=False), document_id="d") if b.bbox)
    f = finding(block, 0, 5) | {"source_map": None}
    stats = annotate_pdf(src, [f], tmp_path / "out.pdf")
    assert stats.placed == 1


@needs_font
@pytest.mark.tesseract
def test_scan_highlight_uses_ocr_words(tmp_path):
    contract = make_contract_pdf(tmp_path / "contract.pdf")
    scan = make_scan(contract, tmp_path / "scan.pdf", style=ScanStyle(max_skew=0.3), pages=[1])
    blocks = build_structure(extract_pdf(scan), document_id="d")
    block = max(blocks, key=lambda b: len(b.text))
    assert block.source_map.words and block.source_map.words[0].confidence is not None
    f = finding(block, 0, min(len(block.text), 30))
    out = tmp_path / "out.pdf"
    assert annotate_pdf(scan, [f], out).placed == 1
    with pymupdf.open(out) as doc:
        page = doc[block.page - 1]
        (annot,) = list(page.annots())
        first = block.source_map.words[0].bbox
        assert annot.rect.intersects(pymupdf.Rect(first.x0, first.y0, first.x1, first.y1))
