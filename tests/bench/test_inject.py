"""Стенд внесения рисков: правка попадает в текст, эталон совпадает с тем, что внесено."""

from __future__ import annotations

from docx import Document

from bench.inject import INJECTIONS, inject_package
from contracts.analysis import FindingCandidate
from ml.extract import extract_document


def _package(tmp_path):
    package = tmp_path / "package"
    package.mkdir()
    contract = Document()
    contract.add_paragraph("Проект контракта на поставку товара")
    contract.add_paragraph("Обеспечение исполнения Контракта устанавливается в размере 5% от начальной цены.")
    contract.add_paragraph("Заказчик оплачивает Поставщику Товар в течение 7 (семи) рабочих дней.")
    contract.save(package / "contract.docx")
    notice = Document()
    notice.add_paragraph("Извещение")
    notice.save(package / "notice.docx")
    return package


def test_injection_changes_text_and_records_expected(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    paths, expected = inject_package(_package(tmp_path), out, only={"R-02"})
    text = "\n".join(p.text for p in extract_document(next(p for p in paths if p.stem == "contract")).paragraphs)
    assert "25 % от начальной" in text
    assert "в размере 5%" not in text  # значение заменено, а не добавлено рядом
    assert "в сумме 15 процентов" in text  # трудная правка — новым абзацем
    assert {e["rule"] for e in expected} == {"R-02"}
    assert all(e["quote"] in e["text"] for e in expected)
    assert {p.stem for p in paths} == {"contract", "notice"}


def test_injection_quotes_are_inside_text():
    for item in INJECTIONS:
        assert item.quote in item.text, item.rule


def test_report_counts_other_part_of_injected_clause_as_hit(tmp_path, capsys):
    from bench.inject import report

    expected = [{"rule": "R-06", "doc": "requirements", "quote": "статус партнёра",
                 "text": "Участник должен иметь статус партнёра и опыт пяти контрактов"}]
    cand = FindingCandidate(rule_code="R-06", level="warning", block_id="b", start=0, end=10,
                            quote="опыт пяти контрактов", explanation="", confidence=0.9, match_score=100)
    row = report("p", {"requirements": [cand]}, expected, [])
    assert row["extra"] == []
